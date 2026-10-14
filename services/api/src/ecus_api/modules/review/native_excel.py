"""Edit formula-free BIFF8 forms while retaining their original layout records."""

import io
import struct
from dataclasses import dataclass
from decimal import Decimal
from typing import Any

import olefile  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from xlwt import BIFFRecords, CompoundDoc  # type: ignore[import-untyped]

from ecus_api.core.errors import AppError
from ecus_api.modules.review.presentation import CellPresentation
from ecus_api.modules.review.schemas import Mapping

CELL_RECORDS = {0x0201, 0x0203, 0x0204, 0x0205, 0x00FD, 0x027E, 0x00BD, 0x00BE}
UNSUPPORTED_RECORDS = {
    0x0006,
    0x0206,
    0x0406,  # Formula caches and their references cannot be relocated safely.
    0x002F,  # Encryption.
    0x005D,
    0x00EC,
    0x00EB,
    0x001C,
    0x01B8,  # Objects, drawings, comments, hyperlinks.
}


@dataclass(frozen=True)
class Record:
    code: int
    data: bytes

    def encode(self) -> bytes:
        return struct.pack("<HH", self.code, len(self.data)) + self.data


@dataclass
class NativeWorkbook:
    book: Any
    globals: list[Record]
    sheets: list[list[Record]]


def read_native(content: bytes) -> NativeWorkbook:
    """Reject features whose cached offsets or hidden content cannot be safely preserved."""
    try:
        with olefile.OleFileIO(io.BytesIO(content)) as archive:
            if any(
                "vba" in part.lower() or part.lower() == "macros"
                for path in archive.listdir()
                for part in path
            ):
                raise ValueError("Macro streams are not supported")
            stream = archive.openstream("Workbook").read()
        book = xlrd.open_workbook(
            file_contents=content, formatting_info=True, logfile=io.StringIO()
        )
        if (
            book.biff_version != 80
            or book.nsheets > 50
            or any(sheet.nrows * sheet.ncols > 200_000 for sheet in book.sheets())
        ):
            raise ValueError("Unsupported workbook version or size")
        sections: list[list[Record]] = [[]]
        offset = 0
        while offset + 4 <= len(stream) and len(sections) <= book.nsheets + 1:
            code, length = struct.unpack_from("<HH", stream, offset)
            end = offset + 4 + length
            if end > len(stream) or code in UNSUPPORTED_RECORDS:
                raise ValueError("Unsupported workbook feature")
            data = stream[offset + 4 : end]
            if code == 0x0809 and struct.unpack_from("<HH", data) != (0x0600, 0x0005):
                if struct.unpack_from("<HH", data) != (0x0600, 0x0010):
                    raise ValueError("Only BIFF8 worksheets are supported")
            if code == 0x01AE and (length != 4 or data[2:] != b"\x01\x04"):
                raise ValueError("External workbooks are not supported")
            sections[-1].append(Record(code, data))
            offset = end
            if code == 0x000A:
                if len(sections) == book.nsheets + 1:
                    break
                sections.append([])
        if len(sections) != book.nsheets + 1 or any(
            section[0].code != 0x0809 or section[-1].code != 0x000A for section in sections
        ):
            raise ValueError("Incomplete workbook stream")
        return NativeWorkbook(book, sections[0], sections[1:])
    except Exception:
        raise AppError(
            "unsupported_legacy_template",
            "Use a formula-free BIFF8 XLS form without encryption, objects or external links.",
            422,
        ) from None


def record_cells(record: Record) -> list[tuple[int, int, int]]:
    """Return row, column and original XF for individual and packed cell records."""
    if record.code not in CELL_RECORDS:
        return []
    row, column = struct.unpack_from("<HH", record.data)
    if record.code in {0x00BE, 0x00BD}:
        last = struct.unpack_from("<H", record.data, len(record.data) - 2)[0]
        step = 2 if record.code == 0x00BE else 6
        return [
            (row, current, struct.unpack_from("<H", record.data, 4 + (current - column) * step)[0])
            for current in range(column, last + 1)
        ]
    return [(row, column, struct.unpack_from("<H", record.data, 4)[0])]


def validate_native_targets(content: bytes, mappings: list[Mapping], maximum: int) -> None:
    if len(content) > maximum:
        raise AppError("invalid_template", "Template exceeds the file size limit.", 422)
    native = read_native(content)
    targets: set[tuple[str, str]] = set()
    physical = {
        sheet.name: {(row, column) for record in records for row, column, _ in record_cells(record)}
        for sheet, records in zip(native.book.sheets(), native.sheets, strict=True)
    }
    for mapping in mappings:
        row, column = coordinate_to_tuple(mapping.cell)
        row -= 1
        column -= 1
        target = mapping.sheet, mapping.cell
        if target in targets or (row, column) not in physical.get(mapping.sheet, set()):
            raise AppError("invalid_mapping", "Use unique existing writable XLS cells.", 422)
        sheet = native.book.sheet_by_name(mapping.sheet)
        for start, end, left, right in sheet.merged_cells:
            if start <= row < end and left <= column < right and (row, column) != (start, left):
                raise AppError("invalid_mapping", "Merged target is not writable.", 422)
        targets.add(target)


def edit_native(
    content: bytes,
    values: dict[tuple[str, str], str | Decimal | None],
    *,
    captions: dict[str, list[str]] | None = None,
    missing: set[tuple[str, str]] | None = None,
    presentation: dict[str, dict[str, CellPresentation]] | None = None,
    omit_print_areas: set[str] | None = None,
) -> bytes:
    """Rebuild cell strings, retaining native formatting, merges, visibility and print setup.

    An explicit caption allowlist removes all historical values, including unused SST entries.
    Compound metadata streams are deliberately excluded from the sanitized copy.
    Sheets open in Normal view rather than showing page-break preview overlays.
    """
    native = read_native(content)
    book = native.book
    if omit_print_areas and any(
        name.name not in {"Print_Area", "Print_Titles"} for name in book.name_obj_list
    ):
        raise AppError(
            "unsupported_legacy_template",
            "Print-area filtering requires a form without user-defined names.",
            422,
        )
    if captions is not None and set(captions) - set(book.sheet_names()):
        raise AppError("invalid_mapping", "Caption allowlist references an absent sheet.", 422)
    sst = BIFFRecords.SharedStringTable("utf-8")
    xfs = [record.data for record in native.globals if record.code == 0x00E0]
    added_xfs: list[bytes] = []
    added_fonts: list[bytes] = []
    added_formats: list[Record] = []
    presentation_xfs: dict[tuple[int, str], int] = {}
    source_fonts = [record.data for record in native.globals if record.code == 0x0031]
    next_format = max([163, *book.format_map]) + 1

    def present(xf_index: int, style: CellPresentation) -> int:
        nonlocal next_format
        key = xf_index, style.model_dump_json()
        if key in presentation_xfs:
            return presentation_xfs[key]
        xf = bytearray(xfs[xf_index] if xf_index < len(xfs) else added_xfs[xf_index - len(xfs)])
        if style.font_name is not None or style.font_size is not None:
            index = struct.unpack_from("<H", xf)[0]
            font = bytearray(source_fonts[index - 1 if index > 4 else index])
            if style.font_size is not None:
                struct.pack_into("<H", font, 0, round(style.font_size * 20))
            if style.font_name is not None:
                font[14:] = bytes((len(style.font_name), 1)) + style.font_name.encode("utf-16le")
            struct.pack_into("<H", xf, 0, len(book.font_list) + len(added_fonts))
            added_fonts.append(bytes(font))
            xf[9] |= 0x08
        if style.wrap is not None:
            xf[6] = (xf[6] & ~0x08) | (0x08 if style.wrap else 0)
        if style.horizontal is not None:
            xf[6] = (xf[6] & ~7) | {"left": 1, "center": 2, "right": 3}[style.horizontal]
        if style.wrap is not None or style.horizontal is not None:
            xf[9] |= 0x10
        if style.number_format is not None:
            data = BIFFRecords.NumberFormatRecord(next_format, style.number_format).get()
            added_formats.append(Record(0x041E, data[4:]))
            struct.pack_into("<H", xf, 2, next_format)
            next_format += 1
            xf[9] |= 0x04
        index = len(xfs) + len(added_xfs)
        presentation_xfs[key] = index
        added_xfs.append(bytes(xf))
        return index

    highlighted_xfs: dict[int, int] = {}
    palette = next((record for record in native.globals if record.code == 0x0092), None)
    palette_data = bytearray(palette.data) if palette else bytearray()
    color_index = 10
    if missing:
        used_colors: set[int] = set()
        for font in book.font_list:
            used_colors.add(font.colour_index)
        for xf in book.xf_list:
            used_colors.update(
                [xf.background.pattern_colour_index, xf.background.background_colour_index]
            )
            for side in ("left", "right", "top", "bottom", "diag"):
                used_colors.add(getattr(xf.border, f"{side}_colour_index"))
        colors = dict(book.colour_map)
        color_index = next((index for index, rgb in colors.items() if rgb == (242, 204, 204)), 0)
        if not color_index:
            color_index = next((index for index in range(8, 64) if index not in used_colors), 0)
            if not color_index:
                raise AppError(
                    "invalid_template", "No unused color is available for missing cells.", 422
                )
            if not palette_data:
                palette_data = bytearray(struct.pack("<H", 56))
                for index in range(8, 64):
                    palette_data.extend(bytes((*colors[index], 0)))
            start = 2 + (color_index - 8) * 4
            palette_data[start : start + 4] = bytes((242, 204, 204, 0))

    def highlight(xf_index: int) -> int:
        if xf_index not in highlighted_xfs:
            xf = bytearray(xfs[xf_index])
            xf[9] |= 0x40  # Apply this cell's fill rather than inheriting it from a style XF.
            border_fill = struct.unpack_from("<I", xf, 14)[0]
            struct.pack_into("<I", xf, 14, (border_fill & 0x03FFFFFF) | (1 << 26))
            struct.pack_into("<H", xf, 18, color_index | (0x40 << 7))
            highlighted_xfs[xf_index] = len(xfs) + len(added_xfs)
            added_xfs.append(bytes(xf))
        return highlighted_xfs[xf_index]

    sheet_streams: list[bytes] = []
    for sheet, records in zip(book.sheets(), native.sheets, strict=True):
        styles = (presentation or {}).get(sheet.name, {})
        row_heights: dict[int, int] = {}
        new_merges: list[tuple[int, int, int, int]] = []
        for address, style in styles.items():
            row, column = coordinate_to_tuple(address)
            if row > sheet.nrows or column > sheet.ncols:
                raise ValueError("Presentation target is outside the template")
            if style.row_height is not None:
                height = round(style.row_height * 20)
                if row - 1 in row_heights and row_heights[row - 1] != height:
                    raise ValueError("Conflicting row presentation heights")
                row_heights[row - 1] = height
            if style.merge_to:
                end_row, end_column = coordinate_to_tuple(style.merge_to)
                area = (row - 1, row, column - 1, end_column)
                if end_row != row or not column <= end_column <= sheet.ncols:
                    raise ValueError("Presentation merges must extend within one row")
                if area not in sheet.merged_cells:
                    if any(
                        sheet.cell_value(row - 1, c) not in (None, "")
                        for c in range(column, end_column)
                    ):
                        raise ValueError("Presentation merge would remove existing content")
                    if any(
                        a < row <= b and left < end_column and column - 1 < right
                        for a, b, left, right in [*sheet.merged_cells, *new_merges]
                    ):
                        raise ValueError("Presentation merge overlaps an existing range")
                    new_merges.append(area)
        allowed = set(captions.get(sheet.name, [])) if captions is not None else None
        highlighted: set[tuple[int, int]] = set()
        for name, coordinate in missing or set():
            if name != sheet.name:
                continue
            row, column = coordinate_to_tuple(coordinate)
            row -= 1
            column -= 1
            highlighted.add((row, column))
            for top, bottom, left, right in sheet.merged_cells:
                if top <= row < bottom and left <= column < right:
                    highlighted.update(
                        (r, c) for r in range(top, bottom) for c in range(left, right)
                    )
        output: list[bytes] = []
        for record in records:
            if record.code == 0x000A and new_merges:
                data = struct.pack("<H", len(new_merges)) + b"".join(
                    struct.pack("<HHHH", top, bottom - 1, left, right - 1)
                    for top, bottom, left, right in new_merges
                )
                output.append(Record(0x00E5, data).encode())
            if record.code == 0x0208:
                row = struct.unpack_from("<H", record.data)[0]
                if row in row_heights:
                    row_data = bytearray(record.data)
                    struct.pack_into("<H", row_data, 6, row_heights[row])
                    row_data[12] |= 0x40
                    record = Record(record.code, bytes(row_data))
            if record.code == 0x023E:  # WINDOW2: bit 11 selects page-break preview.
                flags = struct.unpack_from("<H", record.data)[0]
                output.append(
                    Record(
                        record.code, struct.pack("<H", flags & ~0x0800) + record.data[2:]
                    ).encode()
                )
                if sheet.show_in_page_break_preview:
                    # Restore the cached Normal-view zoom, not the active preview zoom.
                    output.append(
                        Record(
                            0x00A0,
                            struct.pack("<HH", int(sheet.cooked_normal_view_mag_factor), 100),
                        ).encode()
                    )
                continue
            if record.code == 0x00A0 and sheet.show_in_page_break_preview:
                continue
            # Excel accepts sequential cell tables without these optional absolute-offset indexes.
            if record.code in {0x020B, 0x00D7}:
                continue
            cells = record_cells(record)
            if not cells:
                output.append(record.encode())
                continue
            for row, column, xf_index in cells:
                coordinate = f"{get_column_letter(column + 1)}{row + 1}"
                key = sheet.name, coordinate
                cell = sheet.cell(row, column)
                value = values.get(
                    key, cell.value if allowed is None or coordinate in allowed else None
                )
                if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK} and key not in values:
                    value = None
                if (row, column) in highlighted:
                    xf_index = highlight(xf_index)
                if coordinate in styles:
                    xf_index = present(xf_index, styles[coordinate])
                header = struct.pack("<HHH", row, column, xf_index)
                if value is None:
                    replacement = Record(0x0201, header)
                elif isinstance(value, str):
                    replacement = Record(0x00FD, header + struct.pack("<I", sst.add_str(value)))
                else:
                    replacement = Record(0x0203, header + struct.pack("<d", float(value)))
                output.append(replacement.encode())
        sheet_streams.append(b"".join(output))

    globals_output: list[Record | bytes] = []
    skipping_sst = False
    last_xf = max(index for index, record in enumerate(native.globals) if record.code == 0x00E0)
    first_xf = min(index for index, record in enumerate(native.globals) if record.code == 0x00E0)
    last_font = max(index for index, record in enumerate(native.globals) if record.code == 0x0031)
    name_index = 0
    for index, record in enumerate(native.globals):
        if record.code == 0x0018:
            name = book.name_obj_list[name_index]
            name_index += 1
            if (
                name.name == "Print_Area"
                and name.scope >= 0
                and book.sheet_names()[name.scope] in (omit_print_areas or set())
            ):
                continue
        if record.code == 0x00FC:
            globals_output.append(sst.get_biff_record())
            skipping_sst = True
            continue
        if record.code == 0x003C and skipping_sst:
            continue
        skipping_sst = False
        if record.code == 0x00FF:  # ExtSST refers to offsets in the previous string table.
            continue
        if record.code == 0x0092 and missing:
            record = Record(record.code, bytes(palette_data))
        if index == first_xf:
            globals_output.extend(added_formats)
        globals_output.append(record)
        if index == last_font:
            globals_output.extend(Record(0x0031, font) for font in added_fonts)
        if index == last_xf:
            globals_output.extend(Record(0x00E0, xf) for xf in added_xfs)
            if missing and not palette:
                globals_output.append(Record(0x0092, bytes(palette_data)))
    if not any(record.code == 0x00FC for record in native.globals):
        globals_output.insert(-1, sst.get_biff_record())
    globals_size = sum(
        len(item if isinstance(item, bytes) else item.encode()) for item in globals_output
    )
    position = globals_size
    sheet_index = 0
    encoded_globals: list[bytes] = []
    for item in globals_output:
        if isinstance(item, Record):
            if item.code == 0x0085:
                item = Record(item.code, struct.pack("<I", position) + item.data[4:])
                position += len(sheet_streams[sheet_index])
                sheet_index += 1
            encoded_globals.append(item.encode())
        else:
            encoded_globals.append(item)
    output_file = io.BytesIO()
    CompoundDoc.XlsDoc().save(output_file, b"".join(encoded_globals + sheet_streams))
    return output_file.getvalue()
