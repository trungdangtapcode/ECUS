"""Convert an explicitly allowlisted, formula-free XLS layout into a blank XLSX."""

import io
import struct
from datetime import datetime

import olefile  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from ecus_api.core.errors import AppError
from ecus_api.modules.review.template_library import Layout

BORDER_STYLES = {
    0: None,
    1: "thin",
    2: "medium",
    3: "dashed",
    4: "dotted",
    5: "thick",
    6: "double",
    7: "hair",
    8: "mediumDashed",
    9: "dashDot",
    10: "mediumDashDot",
    11: "dashDotDot",
    12: "mediumDashDotDot",
    13: "slantDashDot",
}


def blank_xls_to_xlsx(content: bytes, layout: Layout) -> bytes:
    try:
        with olefile.OleFileIO(io.BytesIO(content)) as archive:
            if any(
                "vba" in part.lower() or part.lower() == "macros"
                for path in archive.listdir()
                for part in path
            ):
                raise ValueError("Macros are not supported")
            stream = archive.openstream("Workbook" if archive.exists("Workbook") else "Book").read()
        offset = 0
        while offset + 4 <= len(stream):
            record, length = struct.unpack_from("<HH", stream, offset)
            if record in {0x0006, 0x0206, 0x0406, 0x002F}:
                raise ValueError("Formulas or encryption are not supported")
            if (
                record == 0x0809
                and length >= 4
                and struct.unpack_from("<H", stream, offset + 6)[0] == 0x0040
            ):
                raise ValueError("Macro sheets are not supported")
            offset += 4 + length
        source = xlrd.open_workbook(
            file_contents=content, formatting_info=True, logfile=io.StringIO()
        )
        if source.nsheets > 50 or any(
            sheet.nrows * sheet.ncols > 200_000 for sheet in source.sheets()
        ):
            raise ValueError("Layout is too large")
        if set(layout.caption_cells) - set(source.sheet_names()):
            raise ValueError("Caption allowlist references an absent sheet")
        workbook = Workbook()
        workbook.remove(workbook.active)  # type: ignore[arg-type]
        workbook.properties.created = datetime(2000, 1, 1)
        workbook.properties.modified = datetime(2000, 1, 1)

        def color(index: int) -> str:
            rgb = source.colour_map.get(index)
            return "FF" + "".join(f"{item:02X}" for item in rgb) if rgb else "FF000000"

        for original in source.sheets():
            sheet = workbook.create_sheet(original.name)
            sheet.sheet_view.showGridLines = False
            allowed = set(layout.caption_cells.get(original.name, []))
            for index, info in original.colinfo_map.items():
                dimension = sheet.column_dimensions[get_column_letter(index + 1)]
                dimension.width = info.width / 256
                dimension.hidden = bool(info.hidden)
            for index, info in original.rowinfo_map.items():
                dimension = sheet.row_dimensions[index + 1]
                dimension.height = info.height / 20
                dimension.hidden = bool(info.hidden)
            for row in range(original.nrows):
                for column in range(original.ncols):
                    old = original.cell(row, column)
                    cell = sheet.cell(row + 1, column + 1)
                    xf = source.xf_list[old.xf_index]
                    font = source.font_list[xf.font_index]
                    cell.font = Font(
                        name=font.name,
                        size=font.height / 20,
                        bold=bool(font.bold),
                        italic=bool(font.italic),
                        color=color(font.colour_index),
                        strike=bool(font.struck_out),
                        underline="single" if font.underlined else None,
                    )
                    alignment = xf.alignment
                    cell.alignment = Alignment(
                        horizontal={
                            1: "left",
                            2: "center",
                            3: "right",
                            5: "justify",
                            6: "centerContinuous",
                            7: "distributed",
                        }.get(alignment.hor_align),
                        vertical={
                            0: "top",
                            1: "center",
                            2: "bottom",
                            3: "justify",
                            4: "distributed",
                        }.get(alignment.vert_align),
                        wrap_text=bool(alignment.text_wrapped),
                        text_rotation=min(alignment.rotation, 180),
                    )
                    border = xf.border
                    sides = {}
                    for side in ("left", "right", "top", "bottom"):
                        sides[side] = Side(
                            style=BORDER_STYLES.get(getattr(border, f"{side}_line_style")),  # type: ignore[arg-type]
                            color=color(getattr(border, f"{side}_colour_index")),
                        )
                    cell.border = Border(
                        left=sides["left"],
                        right=sides["right"],
                        top=sides["top"],
                        bottom=sides["bottom"],
                    )
                    background = xf.background
                    if background.fill_pattern == 1:
                        cell.fill = PatternFill(
                            fill_type="solid", fgColor=color(background.pattern_colour_index)
                        )
                    cell.number_format = source.format_map[xf.format_key].format_str
                    if cell.coordinate in allowed:
                        cell.value = old.value
                        if isinstance(old.value, str):
                            cell.data_type = "s"
            for row_start, row_end, column_start, column_end in original.merged_cells:
                sheet.merge_cells(
                    start_row=row_start + 1,
                    end_row=row_end,
                    start_column=column_start + 1,
                    end_column=column_end,
                )
            for coordinate, title in layout.headings.get(original.name, {}).items():
                sheet[coordinate] = title
            for page_break in layout.page_breaks.get(original.name, []):
                sheet.row_breaks.append(Break(id=page_break))
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.orientation = "portrait"
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.print_area = f"A1:{get_column_letter(original.ncols)}{original.nrows}"
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        source.release_resources()
        return output.getvalue()
    except Exception:
        raise AppError(
            "unsupported_legacy_template",
            "Use a formula-free, unencrypted XLS form with an explicit caption allowlist.",
            422,
        ) from None
