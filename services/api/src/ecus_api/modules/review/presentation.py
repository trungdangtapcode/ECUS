"""Explicit template presentation settings, independent of shipment values."""

import io
from copy import copy
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from pydantic import Field

from ecus_api.modules.review.schemas_base import Contract


class CellPresentation(Contract):
    font_name: str | None = Field(default=None, min_length=1, max_length=31)
    font_size: float | None = Field(default=None, ge=7, le=30)
    wrap: bool | None = None
    horizontal: Literal["left", "center", "right"] | None = None
    row_height: float | None = Field(default=None, ge=10, le=120)
    merge_to: str | None = Field(default=None, pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
    number_format: str | None = Field(default=None, min_length=1, max_length=100)


def apply_presentation(
    content: bytes, file_format: str, settings: dict[str, dict[str, CellPresentation]]
) -> bytes:
    if not settings:
        return content
    if file_format == "xls":
        from ecus_api.modules.review.native_excel import edit_native

        return edit_native(content, {}, presentation=settings)
    book = load_workbook(io.BytesIO(content), keep_links=False)
    for name, cells in settings.items():
        sheet = book[name]
        for address, style in cells.items():
            cell = sheet[address]
            if style.merge_to:
                row, column = coordinate_to_tuple(style.merge_to)
                if row != cell.row or column < cell.column:
                    raise ValueError("Presentation merges must extend within one row")
                if any(
                    sheet.cell(row, c).value is not None for c in range(cell.column + 1, column + 1)
                ):
                    raise ValueError("Presentation merge would remove existing content")
                sheet.merge_cells(f"{address}:{style.merge_to}")
            font, alignment = copy(cell.font), copy(cell.alignment)
            if style.font_name is not None:
                font.name = style.font_name
            if style.font_size is not None:
                font.sz = style.font_size
            if style.wrap is not None:
                alignment.wrap_text = style.wrap
            if style.horizontal is not None:
                alignment.horizontal = style.horizontal
            cell.font, cell.alignment = font, alignment
            if style.row_height is not None:
                sheet.row_dimensions[cell.row].height = style.row_height
            if style.number_format is not None:
                cell.number_format = style.number_format
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
