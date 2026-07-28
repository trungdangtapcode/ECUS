import asyncio
import base64
import binascii
import io
import zipfile
from decimal import Decimal
from uuid import UUID, uuid4

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import coordinate_to_tuple
from pypdf import PdfReader
from sqlalchemy import update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from ecus_api.core.errors import AppError
from ecus_api.db.base import utcnow
from ecus_api.modules.cases.service import get_case
from ecus_api.modules.review.domain import (
    comparison_value,
    effective_pages,
    evaluate,
    normalize,
    normalize_field,
    source_fingerprint,
)
from ecus_api.modules.review.models import CaseReview
from ecus_api.modules.review.schemas import (
    AuditEvent,
    CanonicalInput,
    CanonicalValue,
    Document,
    ExportRecord,
    Observation,
    ObservationInput,
    Selection,
    SelectionInput,
    Template,
    TemplateInput,
    Upload,
    Workspace,
)
from ecus_api.storage.local import LocalBlobStore


async def load(session: AsyncSession, case_id: UUID) -> Workspace:
    await get_case(session, case_id)
    record = await session.get(CaseReview, case_id, populate_existing=True)
    return Workspace.model_validate(record.data) if record else Workspace()


def audit(state: Workspace, action: str, detail: str) -> None:
    state.audit.append(AuditEvent(id=uuid4(), at=utcnow(), action=action, detail=detail))


async def save(
    session: AsyncSession,
    case_id: UUID,
    state: Workspace,
    expected: int,
    *,
    preserve_approval: bool = False,
) -> None:
    if state.revision != expected:
        raise AppError(
            "revision_conflict", "The workspace has changed. Reload before retrying.", 409
        )
    state.revision += 1
    state.approved_revision = state.revision if preserve_approval else None
    data = state.model_dump(mode="json")
    try:
        if expected == 0:
            session.add(CaseReview(case_id=case_id, revision=state.revision, data=data))
            await session.flush()
        else:
            result = await session.execute(
                update(CaseReview)
                .where(CaseReview.case_id == case_id, CaseReview.revision == expected)
                .values(revision=state.revision, data=data)
                .returning(CaseReview.case_id)
            )
            if result.scalar_one_or_none() is None:
                raise AppError(
                    "revision_conflict", "The workspace has changed. Reload before retrying.", 409
                )
        await session.commit()
    except IntegrityError:
        await session.rollback()
        raise AppError(
            "revision_conflict", "The workspace has changed. Reload before retrying.", 409
        ) from None


def decode_upload(data: Upload, maximum: int) -> bytes:
    try:
        content = base64.b64decode(data.content_base64, validate=True)
    except (binascii.Error, ValueError):
        raise AppError("invalid_file", "Invalid file encoding.", 422) from None
    if not content or len(content) > maximum:
        raise AppError("file_size_limit", "File exceeds the configured size limit.", 422)
    return content


def read_pdf(content: bytes, page_limit: int) -> list[str]:
    try:
        if not content.startswith(b"%PDF-"):
            raise ValueError
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or not 0 < len(reader.pages) <= page_limit:
            raise ValueError
        return [(page.extract_text() or "") for page in reader.pages]
    except Exception:
        raise AppError(
            "unsupported_pdf", "The PDF is invalid, encrypted or exceeds the page limit.", 422
        ) from None


async def add_document(
    state: Workspace, data: Upload, store: LocalBlobStore, maximum: int, page_limit: int
) -> None:
    if not data.filename.lower().endswith(".pdf"):
        raise AppError("unsupported_pdf", "This intake currently supports PDF files.", 422)
    content = decode_upload(data, maximum)
    pages = await asyncio.to_thread(read_pdf, content, page_limit)
    key = await asyncio.to_thread(store.put, content)
    if any(document.storage_key == key and not document.withdrawn for document in state.documents):
        raise AppError("duplicate_document", "This document is already in the case.", 409)
    state.documents.append(
        Document(id=uuid4(), filename=data.filename, storage_key=key, pages=pages)
    )
    audit(state, "document_uploaded", data.filename)


def add_observation(state: Workspace, data: ObservationInput) -> Observation:
    document = next(
        (item for item in state.documents if item.id == data.document_id and not item.withdrawn),
        None,
    )
    if document is None or data.page > len(document.pages):
        raise AppError("source_not_found", "The source page does not belong to this case.", 422)
    text = effective_pages(document)[data.page - 1]
    if data.source_offset is None and text.count(data.quote) > 1:
        raise AppError(
            "source_quote_ambiguous", "Select the exact occurrence in the source text.", 422
        )
    offset = text.find(data.quote) if data.source_offset is None else data.source_offset
    if offset < 0 or text[offset : offset + len(data.quote)] != data.quote:
        raise AppError(
            "source_quote_missing", "The quote must occur exactly on the source page.", 422
        )
    if any(
        item.state == "active"
        and item.field == data.field
        and item.document_id == data.document_id
        and item.page == data.page
        and item.quote == data.quote
        and item.offset == offset
        for item in state.observations
    ):
        raise AppError(
            "duplicate_observation", "This field already references this source text.", 409
        )
    rule = (
        next((item for item in state.policy.fields if item.field == data.field), None)
        if state.policy
        else None
    )
    normalized, unit = (
        normalize_field(data.quote, rule, data.unit, data.number_locale)
        if rule and rule.value_type == data.value_type
        else normalize(data.quote, data.value_type, data.unit, data.number_locale)
    )
    observation = Observation(
        source_kind=(
            "human_transcript"
            if data.page in document.transcripts
            else "ocr"
            if data.page in document.ocr_pages
            else "pdf_text"
        ),
        id=uuid4(),
        field=data.field,
        document_id=data.document_id,
        document_type=data.document_type,
        page=data.page,
        quote=data.quote,
        offset=offset,
        normalized=normalized,
        value_type=data.value_type,
        unit=unit,
        verified=data.verified,
        number_locale=data.number_locale,
        source_unit=data.unit,
    )
    state.observations.append(observation)
    state.selections.pop(data.field, None)
    audit(
        state,
        "field_recorded",
        f"{data.field}: {document.filename}, page {data.page}, "
        f"offset {offset}, verified={data.verified}",
    )

    return observation


def select_value(state: Workspace, field: str, data: SelectionInput) -> None:
    observation = next(
        (
            item
            for item in state.observations
            if item.id == data.observation_id and item.field == field and item.state == "active"
        ),
        None,
    )
    if observation is None or not observation.verified:
        raise AppError("source_not_verified", "Select a verified observation for this field.", 422)
    previous = state.selections.get(field)
    if field in state.canonical_values:
        state.canonical_history.append({field: state.canonical_values.pop(field)})
    state.selections[field] = Selection(observation_id=observation.id, reason=data.reason)
    audit(
        state,
        "canonical_value_selected",
        f"{field}: {previous.observation_id if previous else 'unset'} -> "
        f"{observation.id}; {data.reason}",
    )


def verify_observation(state: Workspace, observation_id: UUID, reason: str) -> None:
    observation = next(
        (
            item
            for item in state.observations
            if item.id == observation_id and item.state == "active"
        ),
        None,
    )
    if observation is None:
        raise AppError("source_not_found", "Observation not found.", 404)
    observation.verified = True
    audit(state, "source_verified", f"{observation_id}; {reason}")


def workbook_bytes_valid(content: bytes, maximum: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > maximum * 5:
                raise ValueError
            if any("vbaproject" in item.filename.lower() for item in entries):
                raise ValueError
    except (ValueError, zipfile.BadZipFile):
        raise AppError("invalid_template", "Invalid or oversized XLSX template.", 422) from None


def validate_template(content: bytes, data: TemplateInput, maximum: int) -> None:
    if data.filename.lower().endswith(".xls"):
        from ecus_api.modules.review.native_excel import validate_native_targets

        validate_native_targets(content, data.mappings, maximum)
        return
    workbook_bytes_valid(content, maximum)
    try:
        workbook = load_workbook(io.BytesIO(content), keep_links=False)
        targets: set[tuple[str, str]] = set()
        for mapping in data.mappings:
            row, column = coordinate_to_tuple(mapping.cell)
            if row > 1_048_576 or column > 16_384 or mapping.sheet not in workbook.sheetnames:
                raise ValueError
            target = (mapping.sheet, mapping.cell)
            if target in targets:
                raise ValueError
            cell = workbook[mapping.sheet][mapping.cell]
            if isinstance(cell, MergedCell) or cell.data_type == "f":
                raise ValueError
            targets.add(target)
        workbook.close()
    except Exception:
        raise AppError(
            "invalid_mapping",
            "Use unique writable target cells in the uploaded template. "
            "Formula cells cannot be overwritten.",
            422,
        ) from None


async def set_template(
    state: Workspace, data: TemplateInput, store: LocalBlobStore, maximum: int
) -> None:
    if not data.filename.lower().endswith((".xlsx", ".xls")):
        raise AppError(
            "unsupported_template",
            "Upload an XLSX or supported formula-free BIFF8 XLS template.",
            422,
        )
    content = decode_upload(data, maximum)
    await asyncio.to_thread(validate_template, content, data, maximum)
    key = await asyncio.to_thread(store.put, content)
    for previous in state.template_history:
        if (previous.version == data.version and previous.storage_key != key) or (
            previous.mapping_version == data.mapping_version and previous.mappings != data.mappings
        ):
            raise AppError(
                "version_conflict", "Changed templates or mappings require a new version.", 409
            )
    state.template = Template(
        filename=data.filename,
        storage_key=key,
        version=data.version,
        mapping_version=data.mapping_version,
        mappings=data.mappings,
    )
    state.template_history.append(state.template.model_copy(deep=True))
    audit(
        state,
        "template_registered",
        f"{data.filename}; template={data.version}; mapping={data.mapping_version}; sha256={key}",
    )


def canonical_output(state: Workspace, field: str) -> tuple[str, str] | None:
    canonical = state.canonical_values.get(field)
    if canonical and canonical.source_fingerprint == source_fingerprint(state, field):
        return canonical.value, canonical.value_type
    selection = state.selections.get(field)
    observation = next(
        (
            item
            for item in state.observations
            if selection
            and item.id == selection.observation_id
            and item.state == "active"
            and item.verified
        ),
        None,
    )
    return (observation.normalized, observation.value_type) if observation else None


def edit_canonical(state: Workspace, field: str, data: CanonicalInput) -> None:
    rule = (
        next((rule for rule in state.policy.fields if rule.field == field), None)
        if state.policy
        else None
    )
    if (
        rule is None
        or not state.template
        or field not in {item.field for item in state.template.mappings}
    ):
        raise AppError("field_not_configured", "Choose a mapped field in the active profile.", 422)
    value, unit = normalize_field(data.value, rule, rule.unit, data.number_locale)
    if not value.strip():
        raise AppError("invalid_canonical_value", "A consolidated value cannot be blank.", 422)
    chosen_source = None
    if data.source_observation_id is not None:
        active_documents = {item.id for item in state.documents if not item.withdrawn}
        chosen_source = next(
            (
                item
                for item in state.observations
                if item.id == data.source_observation_id
                and item.field == field
                and item.state == "active"
                and item.document_id in active_documents
            ),
            None,
        )
        if chosen_source is None:
            raise AppError("source_not_found", "Choose an active source for this field.", 422)
        if chosen_source.value_type != rule.value_type or (
            rule.unit and chosen_source.unit != rule.unit
        ):
            raise AppError("field_rule_mismatch", "The source conflicts with the field rule.", 422)
        source_value, source_unit = normalize_field(
            chosen_source.normalized, rule, chosen_source.unit, "en"
        )
        matches_source = (
            Decimal(value) == Decimal(source_value)
            if rule.value_type == "decimal"
            else value == source_value
        )
        if not matches_source:
            raise AppError(
                "source_value_mismatch", "The value does not match the selected source.", 422
            )
        unit = source_unit
        value = source_value
    previous = state.canonical_values.get(field)
    if previous:
        state.canonical_history.append({field: previous.model_copy(deep=True)})
    state.canonical_values[field] = CanonicalValue(
        reviewed=data.reviewed,
        source_observation_id=chosen_source.id if chosen_source else None,
        value=value,
        value_type=rule.value_type,
        unit=unit,
        reason=data.reason,
        source_ids=(
            [chosen_source.id]
            if chosen_source
            else [
                item.id
                for item in state.observations
                if item.field == field and item.state == "active"
            ]
        ),
        source_fingerprint=source_fingerprint(state, field),
        at=utcnow(),
    )
    audit(
        state,
        "canonical_value_selected" if chosen_source else "canonical_value_edited",
        f"{field}; source={chosen_source.id}; {data.reason}"
        if chosen_source
        else f"{field}; {data.reason}",
    )


def confirm_consistent_fields(state: Workspace, fields: list[str], reason: str) -> None:
    """An explicit operator action verifies agreeing evidence and selects its source."""
    from ecus_api.modules.review.schemas import Selection

    active_documents = {item.id for item in state.documents if not item.withdrawn}
    for field in dict.fromkeys(fields):
        rule = (
            next((item for item in state.policy.fields if item.field == field), None)
            if state.policy
            else None
        )
        sources = [
            item
            for item in state.observations
            if item.field == field
            and item.state == "active"
            and item.document_id in active_documents
        ]
        if not rule or not sources or field in state.canonical_values:
            raise AppError(
                "consistent_review_unavailable", "Choose supported, unedited source fields.", 422
            )
        values = {comparison_value(item, rule) for item in sources}
        if len(values) != 1 or any(
            item.value_type != rule.value_type or (rule.unit and item.unit != rule.unit)
            for item in sources
        ):
            raise AppError(
                "consistent_review_unavailable",
                "Conflicting fields require individual review.",
                422,
            )
        preferred = sorted(
            [item for item in sources if item.document_type in rule.preferred_sources],
            key=lambda item: rule.preferred_sources.index(item.document_type),
        )
        if rule.preferred_sources and not preferred:
            raise AppError(
                "consistent_review_unavailable",
                "The preferred source is absent. Review an exception.",
                422,
            )
        chosen = (preferred or sources)[0]
        classifications = [
            item
            for item in state.logical_documents
            if not item.withdrawn
            and item.document_id == chosen.document_id
            and item.document_type == chosen.document_type
            and item.page_from <= chosen.page <= item.page_to
        ]
        if not classifications:
            raise AppError(
                "source_classification_required", "Confirm the source classification first.", 422
            )
        for source in sources:
            source.verified = True
            audit(state, "field_verified", f"{source.id}; {reason}")
        for classification in classifications:
            classification.verified = True
        state.selections[field] = Selection(observation_id=chosen.id, reason=reason)
        audit(state, "canonical_value_selected", f"{field}; {chosen.id}; {reason}")


def generate_excel(
    state: Workspace,
    content: bytes,
    output_values: dict[str, tuple[str, str] | None] | None = None,
) -> bytes:
    if state.template is None:
        raise AppError("export_not_ready", "No template registered.", 409)
    if state.template.filename.lower().endswith(".xls"):
        from ecus_api.modules.review.native_excel import edit_native

        values: dict[tuple[str, str], str | Decimal | None] = {}
        missing: set[tuple[str, str]] = set()
        for mapping in state.template.mappings:
            value = (
                output_values.get(mapping.field)
                if output_values is not None
                else canonical_output(state, mapping.field)
            )
            key = mapping.sheet, mapping.cell
            values[key] = (
                None if value is None else Decimal(value[0]) if value[1] == "decimal" else value[0]
            )
            if value is None:
                missing.add(key)
        return edit_native(content, values, missing=missing)
    workbook = load_workbook(io.BytesIO(content), keep_links=False)
    red = PatternFill(fill_type="solid", fgColor="FFF2CCCC")
    for mapping in state.template.mappings:
        value = (
            output_values.get(mapping.field)
            if output_values is not None
            else canonical_output(state, mapping.field)
        )
        sheet = workbook[mapping.sheet]
        cell = sheet[mapping.cell]
        if isinstance(cell, MergedCell):
            raise AppError("invalid_mapping", "Merged target is not writable.", 422)
        if value is None:
            cell.value = None
            cell.fill = red
            for merged in sheet.merged_cells.ranges:
                if mapping.cell in merged:
                    for row in sheet.iter_rows(
                        min_row=merged.min_row,
                        max_row=merged.max_row,
                        min_col=merged.min_col,
                        max_col=merged.max_col,
                    ):
                        for target in row:
                            target.fill = red
        elif value[1] == "decimal":
            cell.value = Decimal(value[0])
        else:
            cell.value = value[0]
            cell.data_type = "s"  # Preserve identifiers and keep formula-like text literal.
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


async def export_excel(
    state: Workspace,
    store: LocalBlobStore,
    *,
    draft: bool = False,
    reason: str = "",
) -> ExportRecord:
    readiness = evaluate(state)
    if not draft and not readiness.approved:
        raise AppError(
            "export_not_ready",
            "The current revision must be reviewed and approved by the server.",
            409,
        )
    if state.template is None:
        raise AppError("export_not_ready", "Register a template before exporting.", 409)
    resolved = {item.field for item in readiness.comparisons if item.resolved}
    values = {
        mapping.field: (
            canonical_output(state, mapping.field)
            if not draft or mapping.field in resolved
            else None
        )
        for mapping in state.template.mappings
    }
    content = await asyncio.to_thread(store.read, state.template.storage_key)
    output = await asyncio.to_thread(generate_excel, state, content, values)
    key = await asyncio.to_thread(store.put, output)
    record = ExportRecord(
        kind="draft" if draft else "approved",
        blockers_snapshot=list(readiness.blockers),
        annotation_fields=[],
        reason=reason,
        observations_snapshot=[item.model_copy(deep=True) for item in state.observations],
        ocr_snapshot={
            document.id: document.ocr_pages for document in state.documents if document.ocr_pages
        },
        transcript_snapshot={
            document.id: dict(document.transcripts)
            for document in state.documents
            if document.transcripts
        },
        canonical_snapshot={
            field: item.model_copy(deep=True) for field, item in state.canonical_values.items()
        },
        values_snapshot={field: item[0] if item else None for field, item in values.items()},
        missing_fields=[field for field, item in values.items() if item is None],
        policy_snapshot=state.policy,
        stage=state.stage,
        applicability_snapshot=state.applicability,
        logical_documents_snapshot=state.logical_documents,
        mapping_snapshot=state.template.mappings,
        selected_observations={
            field: selected.observation_id for field, selected in state.selections.items()
        },
        id=uuid4(),
        at=utcnow(),
        filename=f"{'draft-review' if draft else 'review'}-{state.revision}."
        f"{'xls' if state.template.filename.lower().endswith('.xls') else 'xlsx'}",
        storage_key=key,
        source_revision=state.revision,
        template_version=state.template.version,
        mapping_version=state.template.mapping_version,
    )
    state.exports.append(record)
    audit(
        state,
        "excel_draft_exported" if draft else "excel_exported",
        f"{record.id}; revision={state.revision}; template={record.template_version}; "
        f"mapping={record.mapping_version}; {reason}",
    )
    return record


async def export_draft(state: Workspace, store: LocalBlobStore, reason: str) -> ExportRecord:
    return await export_excel(state, store, draft=True, reason=reason)
