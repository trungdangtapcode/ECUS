"""Database-backed blank layouts and their explicit, versioned mapping contracts."""

import asyncio
import base64
import hashlib
import io
import json
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from pydantic import Field, model_validator
from sqlalchemy import JSON, String, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import Mapped, mapped_column

from ecus_api.core.errors import AppError
from ecus_api.db.base import Base
from ecus_api.modules.review import revisions, service
from ecus_api.modules.review.policy import ArithmeticRule, FieldRule, ReviewPolicy
from ecus_api.modules.review.presentation import CellPresentation
from ecus_api.modules.review.schemas import (
    Mapping,
    PolicyInput,
    TemplateInput,
    VerifyInput,
    Workspace,
)
from ecus_api.modules.review.schemas_base import Contract
from ecus_api.storage.local import LocalBlobStore


class TemplateProfile(Base):
    __tablename__ = "review_template_profiles"
    id: Mapped[str] = mapped_column(String(80), primary_key=True)
    definition: Mapped[dict[str, Any]] = mapped_column(JSON)
    storage_key: Mapped[str] = mapped_column(String(64))
    enabled: Mapped[bool] = mapped_column(default=True)


class Target(Contract):
    sheet: str
    cell: str = Field(pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")


class LayoutField(FieldRule):
    targets: list[Target] = Field(min_length=1)
    required: bool = False


class Slot(Contract):
    sheet: str
    row_offset: int = Field(default=0, ge=0)
    column_offset: int = Field(default=0, ge=0)
    field_targets: dict[str, list[Target]] = Field(default_factory=dict)


class RepeatGroup(Contract):
    arithmetic_checks: list[ArithmeticRule] = Field(default_factory=list)
    key: str = Field(pattern=r"^[a-z][a-z0-9_]{0,20}$")
    maximum: int = Field(ge=1, le=100)
    fields: list[LayoutField] = Field(min_length=1)
    slots: list[Slot] = Field(min_length=1)
    repeat_sheet: str | None = None
    repeat_offsets: list[int] = Field(default_factory=list)


class Layout(Contract):
    presentation: dict[str, dict[str, CellPresentation]] = Field(default_factory=dict)
    continuation_sheets: list[str] = Field(default_factory=list)
    id: str = Field(pattern=r"^[a-zA-Z0-9_-]{1,80}$")
    name_key: str
    version: str = Field(pattern=r"^[a-zA-Z0-9_-]{1,50}$")
    file_format: Literal["xlsx", "xls"] = "xlsx"
    caption_cells: dict[str, list[str]]
    headings: dict[str, dict[str, str]] = Field(default_factory=dict)
    page_breaks: dict[str, list[int]] = Field(default_factory=dict)
    fields: list[LayoutField] = Field(min_length=1)
    groups: list[RepeatGroup] = Field(default_factory=list)

    @model_validator(mode="after")
    def unique_identifiers(self) -> "Layout":
        if len({item.field for item in self.fields}) != len(self.fields) or len(
            {item.key for item in self.groups}
        ) != len(self.groups):
            raise ValueError("Layout identifiers must be unique")
        for group in self.groups:
            if self.file_format == "xls" and group.repeat_sheet:
                raise ValueError("Native XLS layouts require explicit fixed slots")
            if len({item.field for item in group.fields}) != len(group.fields):
                raise ValueError("Repeated field identifiers must be unique")
            for slot in group.slots:
                if set(slot.field_targets) - {item.field for item in group.fields} or any(
                    not targets for targets in slot.field_targets.values()
                ):
                    raise ValueError("Slot overrides must name configured fields and targets")
            if group.repeat_sheet and not group.repeat_offsets:
                raise ValueError("Repeated sheets require explicit block offsets")
            if not group.repeat_sheet and group.maximum > len(group.slots):
                raise ValueError("Fixed layouts cannot silently truncate repeated records")
        return self


class TemplateProfileRead(Contract):
    id: str
    name_key: str
    version: str
    groups: dict[str, int]


class ProfileInput(VerifyInput):
    profile_id: str
    adopt_profile_rules: bool = False
    counts: dict[str, int] = Field(default_factory=dict)


async def list_profiles(session: AsyncSession) -> list[TemplateProfileRead]:
    records = await session.scalars(
        select(TemplateProfile)
        .where(TemplateProfile.enabled.is_(True))
        .order_by(TemplateProfile.id.desc())
    )
    result = []
    for record in records:
        layout = Layout.model_validate(record.definition)
        result.append(
            TemplateProfileRead(
                id=record.id,
                name_key=layout.name_key,
                version=layout.version,
                groups={group.key: group.maximum for group in layout.groups},
            )
        )
    return result


def expand_layout(
    content: bytes, layout: Layout, counts: dict[str, int]
) -> tuple[bytes, list[Mapping], list[FieldRule]]:
    if set(counts) - {group.key for group in layout.groups}:
        raise AppError("invalid_repeat_count", "Unknown repeated record group.", 422)
    workbook = (
        load_workbook(io.BytesIO(content), keep_links=False)
        if layout.file_format == "xlsx"
        else None
    )
    mappings: list[Mapping] = []
    rules: list[FieldRule] = []
    cloned_sheets: dict[str, list[str]] = {}

    def append(item: LayoutField, field: str, slot: Slot | None = None) -> None:
        rules.append(FieldRule(**{**item.model_dump(exclude={"targets"}), "field": field}))
        overrides = slot.field_targets.get(item.field) if slot else None
        for target in overrides or item.targets:
            row, column = coordinate_to_tuple(target.cell)
            target_row = row + (slot.row_offset if slot and not overrides else 0)
            target_column = column + (slot.column_offset if slot and not overrides else 0)
            mappings.append(
                Mapping(
                    field=field,
                    sheet=slot.sheet if slot and not overrides else target.sheet,
                    cell=f"{get_column_letter(target_column)}{target_row}",
                    required=item.required,
                )
            )

    for group in layout.groups:
        count = counts.get(group.key, 1)
        if isinstance(count, bool) or not 0 <= count <= group.maximum:
            raise AppError("invalid_repeat_count", "Record count exceeds the layout capacity.", 422)
        slots = list(group.slots)
        while len(slots) < count:
            assert workbook is not None
            assert group.repeat_sheet is not None
            copies = cloned_sheets.setdefault(group.repeat_sheet, [])
            sheet = workbook.copy_worksheet(workbook[group.repeat_sheet])
            sheet.title = f"{group.repeat_sheet}_{len(copies) + 2:02}"
            sheet.print_area = workbook[group.repeat_sheet].print_area
            copies.append(sheet.title)
            slots.extend(
                Slot(sheet=sheet.title, row_offset=offset) for offset in group.repeat_offsets
            )
        for index, slot in enumerate(slots[:count], 1):
            for item in group.fields:
                append(item, f"{group.key}_{index:03}_{item.field}", slot)
    for item in layout.fields:
        append(item, item.field)
        for target in item.targets:
            for clone in cloned_sheets.get(target.sheet, []):
                mappings.append(
                    Mapping(field=item.field, sheet=clone, cell=target.cell, required=item.required)
                )
    if layout.file_format == "xls" and any(
        item.cell in layout.caption_cells.get(item.sheet, []) for item in mappings
    ):
        raise AppError("invalid_mapping", "Data targets must not overwrite retained captions.", 422)
    if workbook is not None:
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        content = output.getvalue()
    elif layout.continuation_sheets:
        from ecus_api.modules.review.native_excel import edit_native

        active = {
            slot.sheet
            for group in layout.groups
            for slot in group.slots[: counts.get(group.key, 1)]
        }
        content = edit_native(
            content, {}, omit_print_areas=set(layout.continuation_sheets) - active
        )
    return content, mappings, rules


async def apply_profile(
    state: Workspace, data: ProfileInput, session: AsyncSession, store: LocalBlobStore, maximum: int
) -> None:
    record = await session.get(TemplateProfile, data.profile_id)
    if record is None or not record.enabled:
        raise AppError(
            "template_profile_missing", "The blank template profile is not available.", 404
        )
    layout = Layout.model_validate(record.definition)
    content = await asyncio.to_thread(store.read, record.storage_key)
    content, mappings, rules = await asyncio.to_thread(expand_layout, content, layout, data.counts)
    configured = {item.field for item in rules}
    prior_policy = state.policy
    previous_mapped = {item.field for item in state.template.mappings} if state.template else set()
    if prior_policy:
        prior_rules = {rule.field: rule for rule in prior_policy.fields}
        if not data.adopt_profile_rules:
            rules = [prior_rules.get(rule.field, rule) for rule in rules]
        rules.extend(
            rule
            for rule in prior_policy.fields
            if rule.field not in configured and rule.field not in previous_mapped
        )

    active = {item.field for item in state.observations if item.state == "active"} | set(
        state.canonical_values
    )
    if active - {rule.field for rule in rules}:
        raise AppError(
            "profile_fields_in_use",
            "The new profile would omit reviewed fields. Withdraw them or retain the record count.",
            409,
        )
    suffix = "-".join(f"{group.key}{data.counts.get(group.key, 1)}" for group in layout.groups)
    version = f"{layout.version}-{suffix}".rstrip("-")
    previous = next(
        (
            item
            for item in state.template_history
            if item.version == version
            and item.mapping_version == version
            and item.mappings == mappings
        ),
        None,
    )
    if previous:
        state.template = previous.model_copy(deep=True)
    else:
        await service.set_template(
            state,
            TemplateInput(
                revision=data.revision,
                filename=f"{layout.id}.{layout.file_format}",
                content_base64=base64.b64encode(content).decode(),
                version=version,
                mapping_version=version,
                mappings=mappings,
            ),
            store,
            maximum,
        )
    assert state.template is not None
    state.template.missing_values = "highlight"
    state.template.profile_id = record.id
    state.template.repeat_counts = {
        group.key: data.counts.get(group.key, 1) for group in layout.groups
    }
    if not previous:
        state.template_history[-1] = state.template.model_copy(deep=True)
    policy_values: dict[str, Any] = (
        prior_policy.model_dump(mode="json", exclude={"version", "fields"})
        if prior_policy
        else {
            "name": layout.name_key,
            "status": "project_baseline",
        }
    )
    policy_values["fields"] = [rule.model_dump() for rule in rules]
    generated_checks = [
        ArithmeticRule(
            **{
                **check.model_dump(),
                "key": f"{group.key}_{index:03}_{check.key}",
                "factors": [f"{group.key}_{index:03}_{field}" for field in check.factors],
                "result_field": f"{group.key}_{index:03}_{check.result_field}",
            }
        )
        for group in layout.groups
        for index in range(1, data.counts.get(group.key, 1) + 1)
        for check in group.arithmetic_checks
    ]
    generated_keys = {check.key for check in generated_checks}
    configured = {rule.field for rule in rules}
    retained_checks = (
        [
            check
            for check in prior_policy.arithmetic_checks
            if prior_policy
            if check.key not in generated_keys
            and set([*check.factors, check.result_field]) <= configured
        ]
        if prior_policy
        else []
    )
    policy_values["arithmetic_checks"] = [
        check.model_dump(mode="json") for check in [*retained_checks, *generated_checks]
    ]
    policy_hash = hashlib.sha256(json.dumps(policy_values, sort_keys=True).encode()).hexdigest()[
        :12
    ]
    revisions.adopt_policy(
        state,
        PolicyInput(
            revision=data.revision,
            reason=data.reason,
            policy=ReviewPolicy.model_validate(
                {**policy_values, "version": f"{layout.version}-{policy_hash}"}
            ),
        ),
    )
    service.audit(state, "template_profile_applied", f"{record.id}; {version}; {data.reason}")
