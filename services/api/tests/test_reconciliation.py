"""Single-call semantic review, hostile-output guards and explicit operator confirmation."""

import asyncio
import base64
import io
import json
from uuid import UUID, uuid4

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from test_review import new_case, pdf_bytes, upload

from ecus_api.ai import tracing
from ecus_api.ai.contracts import AIResult
from ecus_api.ai.openai_compatible import OpenAICompatible
from ecus_api.core.config import Settings
from ecus_api.core.errors import AppError
from ecus_api.modules.review.reconciliation import ReconciliationOutput

VALUES = {
    "package_quantity": [("1250", "BOXES"), ("1250", "BOX")],
    "package_unit": [("Boxes", ""), ("BOX", "")],
    "weight_unit": [("Kgs", ""), ("KG", "")],
    "exporter_phone": [("+84 28 12345678", ""), ("(+84) 28 12345678", "")],
    "company_name": [("Example Company JSC", ""), ("EXAMPLE COMPANY JOINT STOCK COMPANY", "")],
    "total_fees": [("500000", "VND"), ("87000", "VND")],
    "loading_port_name": [("REGION ALPHA, COUNTRY", ""), ("TERMINAL BETA", "")],
}


async def configured(client: httpx.AsyncClient) -> tuple[str, dict]:
    root = await new_case(client)
    rules = [
        {
            "field": field,
            "label": field,
            "value_type": "decimal" if field in {"package_quantity", "total_fees"} else "text",
        }
        for field in VALUES
    ]
    response = await client.post(
        f"{root}/policy",
        json={
            "revision": 0,
            "reason": "Synthetic review policy",
            "policy": {
                "version": "semantic-v1",
                "name": "Synthetic semantic review",
                "status": "project_baseline",
                "fields": rules,
            },
        },
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    book = Workbook()
    book.active.title = "Review"
    book.active["A1"] = "Synthetic semantic review"
    output = io.BytesIO()
    book.save(output)
    book.close()
    response = await client.post(
        f"{root}/template",
        json={
            "revision": state["revision"],
            "filename": "synthetic.xlsx",
            "content_base64": base64.b64encode(output.getvalue()).decode(),
            "version": "semantic-v1",
            "mapping_version": "semantic-v1",
            "mappings": [
                {"field": field, "sheet": "Review", "cell": f"B{index + 2}", "required": False}
                for index, field in enumerate(VALUES)
            ],
        },
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    for index in range(2):
        content = pdf_bytes(
            ["Synthetic source", *[values[index][0] for values in VALUES.values()], "CRATE"]
        )
        response = await client.post(
            f"{root}/documents", json=upload(content, f"source-{index}.pdf", state["revision"])
        )
        assert response.status_code == 200
        state = response.json()["workspace"]
        document_id = state["documents"][-1]["id"]
        for field, values in VALUES.items():
            value, unit = values[index]
            response = await client.post(
                f"{root}/observations",
                json={
                    "revision": state["revision"],
                    "document_id": document_id,
                    "document_type": "INVOICE",
                    "field": field,
                    "quote": value,
                    "page": 1,
                    "unit": unit,
                    "value_type": "decimal"
                    if field in {"package_quantity", "total_fees"}
                    else "text",
                    "verified": False,
                },
            )
            assert response.status_code == 200, response.text
            state = response.json()["workspace"]
    return root, state


def decisions(prompt: dict) -> list[dict]:
    result = []
    for field in prompt["fields"]:
        source = field["sources"][0]
        different = field["field"] in {"total_fees", "loading_port_name"}
        standardized = field["field"] == "company_name"
        result.append(
            {
                "field": field["field"],
                "classification": "different" if different else "equivalent",
                "source_ids": [item["id"] for item in field["sources"]],
                "source_observation_id": None if different or standardized else source["id"],
                "proposed_value": None
                if different
                else "Example Company, Joint Stock Company"
                if standardized
                else source["value"],
                "reason": "Synthetic evidence supports this comparison.",
            }
        )
    return result


def fake_result(schema, prompt: dict):
    return AIResult(
        value=schema.model_validate({"decisions": decisions(prompt)}),
        model="mock-semantic",
        usage={"total_tokens": 1},
        request_id="synthetic-request",
        prompt_version="v1",
    )


async def finished(client: httpx.AsyncClient, root: str) -> dict:
    for _ in range(300):
        run = (await client.get(f"{root}/reconciliations")).json()[0]
        if run["status"] not in {"pending", "running"}:
            return run
        await asyncio.sleep(0.01)
    pytest.fail("Semantic review did not complete")


async def test_one_call_caches_across_concurrency_locale_confirmation_and_export(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    root, state = await configured(client)
    calls = []
    release = asyncio.Event()

    async def fake(_self, text, schema):
        prompt = json.loads(text)
        calls.append(prompt)
        await release.wait()
        return fake_result(schema, prompt)

    monkeypatch.setattr(OpenAICompatible, "reconcile_values", fake)
    original = state["observations"]
    first, second = await asyncio.gather(
        *[
            client.post(
                f"{root}/reconciliations", json={"revision": state["revision"], "locale": "vi"}
            )
            for _ in range(2)
        ]
    )
    assert first.status_code == second.status_code == 202
    assert first.json()["id"] == second.json()["id"]
    release.set()
    run = await finished(client, root)
    assert run["status"] == "succeeded", run
    assert run["request_count"] == 1 and len(calls) == 1
    assert len(calls[0]["fields"]) == len(VALUES) and "pages" not in calls[0]
    data = (await client.get(root)).json()
    state = data["workspace"]
    assert state["observations"] == original and not state["canonical_values"]
    assert not data["readiness"]["approved"] and not data["readiness"]["eligible"]
    assert state["approved_revision"] is None
    comparisons = {item["field"]: item for item in data["readiness"]["comparisons"]}
    assert not comparisons["package_quantity"]["source_values_differ"]
    assert comparisons["total_fees"]["source_values_differ"]
    assert comparisons["loading_port_name"]["source_values_differ"]
    equivalent = [
        item["field"]
        for item in run["batch"]["decisions"]
        if item["classification"] == "equivalent"
    ]
    response = await client.post(
        f"{root}/reconciliations/{run['id']}/apply",
        json={
            "revision": state["revision"],
            "fields": equivalent,
            "reviewed": False,
            "reason": "Not confirmed",
        },
    )
    assert response.status_code == 422
    response = await client.post(
        f"{root}/reconciliations/{run['id']}/apply",
        json={
            "revision": state["revision"],
            "fields": [equivalent[0], "total_fees"],
            "reviewed": True,
            "reason": "Reject partial batch",
        },
    )
    assert response.status_code == 409
    assert (await client.get(root)).json()["workspace"] == state
    response = await client.post(
        f"{root}/reconciliations/{run['id']}/apply",
        json={
            "revision": state["revision"],
            "fields": equivalent,
            "reviewed": True,
            "reason": "Checked the synthetic evidence",
        },
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    assert state["observations"] == original
    assert (
        state["canonical_values"]["company_name"]["value"] == "Example Company, Joint Stock Company"
    )
    assert state["canonical_values"]["company_name"]["source_observation_id"] is None
    assert len(state["canonical_values"]["company_name"]["source_ids"]) == 2
    assert len(state["canonical_values"]["package_quantity"]["source_ids"]) == 1
    assert all(
        item["ai_reconciliation_id"] == run["id"] for item in state["canonical_values"].values()
    )
    cached = await client.post(
        f"{root}/reconciliations", json={"revision": state["revision"], "locale": "en"}
    )
    assert cached.json()["id"] == run["id"] and len(calls) == 1
    exported = await client.post(
        f"{root}/draft-exports",
        json={"revision": state["revision"], "reason": "Review synthetic output"},
    )
    assert exported.status_code == 200, exported.text
    state = exported.json()["workspace"]
    record = state["exports"][-1]
    assert record["ai_reconciliation_snapshot"]["id"] == run["id"]
    content = (await client.get(f"{root}/exports/{record['id']}/content")).content
    book = load_workbook(io.BytesIO(content))
    assert book["Review"]["B2"].value == 1250
    assert book["Review"]["B6"].value == "Example Company, Joint Stock Company"
    assert book["Review"]["B7"].value is None
    book.close()
    assert len(calls) == 1


@pytest.mark.parametrize(
    "scenario",
    [
        "numeric",
        "unit",
        "critical",
        "coverage",
        "foreign_source",
        "invented_numeric",
        "omitted",
        "duplicate",
    ],
)
async def test_guard_rejects_unsafe_or_incomplete_model_equivalence(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch, scenario: str
) -> None:
    root, state = await configured(client)
    target = (
        "total_fees"
        if scenario in {"numeric", "invented_numeric"}
        else "loading_port_name"
        if scenario == "critical"
        else "package_quantity"
    )

    async def fake(_self, text, schema):
        prompt = json.loads(text)
        items = decisions(prompt)
        item = next(item for item in items if item["field"] == target)
        source = next(field for field in prompt["fields"] if field["field"] == target)["sources"][0]
        item.update(
            classification="equivalent",
            proposed_value=source["value"],
            source_observation_id=source["id"],
        )
        if scenario == "coverage":
            item["source_ids"] = item["source_ids"][:1]
        elif scenario == "foreign_source":
            item["source_observation_id"] = str(uuid4())
        elif scenario == "invented_numeric":
            item.update(proposed_value="587000", source_observation_id=None)
        elif scenario == "omitted":
            items.remove(item)
        elif scenario == "duplicate":
            items.append(item.copy())
        elif scenario == "unit":
            # A mismatched proposal cannot redefine the source unit through the response.
            item["proposed_value"] = "1250 KG"
        return AIResult(
            value=schema.model_validate({"decisions": items}),
            model="mock",
            usage={},
            request_id=None,
            prompt_version="v1",
        )

    monkeypatch.setattr(OpenAICompatible, "reconcile_values", fake)
    assert (
        await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    ).status_code == 202
    run = await finished(client, root)
    assert run["status"] == "succeeded", run
    decision = next(item for item in run["batch"]["decisions"] if item["field"] == target)
    assert decision["classification"] == "uncertain" and decision["guard_code"]
    assert decision["proposed_value"] is None
    result = (await client.get(root)).json()
    assert next(item for item in result["readiness"]["comparisons"] if item["field"] == target)[
        "source_values_differ"
    ]
    assert not result["workspace"]["canonical_values"]


async def test_changed_evidence_during_call_retains_result_but_never_applies_it(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    root, state = await configured(client)
    started = asyncio.Event()
    release = asyncio.Event()
    calls = []

    async def fake(_self, text, schema):
        prompt = json.loads(text)
        calls.append(prompt)
        started.set()
        await release.wait()
        return fake_result(schema, prompt)

    monkeypatch.setattr(OpenAICompatible, "reconcile_values", fake)
    await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    await started.wait()
    source = state["observations"][0]
    changed = await client.post(
        f"{root}/observations/{source['id']}/withdraw",
        json={"revision": state["revision"], "reason": "Changed synthetic evidence"},
    )
    assert changed.status_code == 200
    release.set()
    run = await finished(client, root)
    assert run["status"] == "succeeded" and not run["current"] and run["batch"]
    result = (await client.get(root)).json()["workspace"]
    assert (
        result["ai_reconciliation"] is None and not result["canonical_values"] and len(calls) == 1
    )


async def test_existing_operator_choice_is_not_overwritten(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    root, state = await configured(client)

    async def fake(_self, text, schema):
        return fake_result(schema, json.loads(text))

    monkeypatch.setattr(OpenAICompatible, "reconcile_values", fake)
    manual = await client.post(
        f"{root}/fields/company_name/canonical",
        json={
            "revision": state["revision"],
            "value": "Operator supplied company",
            "reviewed": True,
            "reason": "Reviewed manual text",
        },
    )
    state = manual.json()["workspace"]
    await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    run = await finished(client, root)
    state = (await client.get(root)).json()["workspace"]
    response = await client.post(
        f"{root}/reconciliations/{run['id']}/apply",
        json={
            "revision": state["revision"],
            "fields": ["package_quantity", "company_name"],
            "reviewed": True,
            "reason": "Must not overwrite manual choice",
        },
    )
    assert (
        response.status_code == 409
        and response.json()["error"]["code"] == "reconciliation_value_already_reviewed"
    )
    assert (await client.get(root)).json()["workspace"] == state


async def test_failed_request_is_not_replayed_and_oversized_input_makes_no_call(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch, settings: Settings
) -> None:
    root, state = await configured(client)
    calls = []

    async def fail(_self, text, schema):
        calls.append(text)
        raise AppError("ai_timeout", "Synthetic timeout", 504)

    monkeypatch.setattr(OpenAICompatible, "reconcile_values", fail)
    settings.ai_max_input_chars = 10
    response = await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    assert response.status_code == 422 and not calls
    settings.ai_max_input_chars = 100000
    first = await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    run = await finished(client, root)
    assert run["status"] == "failed" and run["error_code"] == "ai_timeout"
    second = await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    assert second.json()["id"] == first.json()["id"] and len(calls) == 1


@pytest.mark.parametrize("dedicated", [False, True])
async def test_compatible_semantic_transport_uses_its_own_prompt_and_private_trace(
    settings: Settings, dedicated: bool
) -> None:
    settings.ai_trace_enabled = True
    settings.ai_reconciliation_model = "semantic-test-model" if dedicated else ""
    calls = []
    case_id = uuid4()

    def respond(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        calls.append(body)
        assert "Review all supplied" in body["messages"][0]["content"]
        assert body["model"] == (
            "semantic-test-model" if dedicated else settings.ai_extraction_model
        )
        return httpx.Response(
            200,
            json={
                "model": body["model"],
                "choices": [{"finish_reason": "stop", "message": {"content": '{"decisions":[]}'}}],
            },
            headers={"x-request-id": "synthetic-semantic-request"},
        )

    with tracing.scope(case_id=case_id, phase="reconciliation", fields=["package_unit"]):
        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
            result = await OpenAICompatible(settings, client).reconcile_values(
                '{"task":"reconcile_source_values","fields":[]}', ReconciliationOutput
            )
    assert len(calls) == 1 and result.trace_id
    entries = tracing.summaries(settings, case_id)
    assert len(entries) == 1 and entries[0]["phase"] == "reconciliation"
    content = tracing.content(settings, case_id, UUID(result.trace_id)).decode()
    assert "test-key" not in content and "synthetic-semantic-request" in content


@pytest.mark.parametrize("preapproved", [False, True])
async def test_semantic_information_preserves_but_never_grants_approval(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch, preapproved: bool
) -> None:
    root, state = await configured(client)
    for field, values in VALUES.items():
        response = await client.post(
            f"{root}/fields/{field}/canonical",
            json={
                "revision": state["revision"],
                "value": values[0][0],
                "reviewed": True,
                "reason": "Operator reviewed the synthetic value",
            },
        )
        assert response.status_code == 200
        state = response.json()["workspace"]
    assert response.json()["readiness"]["eligible"]
    if preapproved:
        response = await client.post(
            f"{root}/approve",
            json={"revision": state["revision"], "reason": "Explicit operator approval"},
        )
        assert response.status_code == 200
        state = response.json()["workspace"]
    prior_values = state["canonical_values"]

    async def fake(_self, text, schema):
        return fake_result(schema, json.loads(text))

    monkeypatch.setattr(OpenAICompatible, "reconcile_values", fake)
    await client.post(f"{root}/reconciliations", json={"revision": state["revision"]})
    run = await finished(client, root)
    assert run["status"] == "succeeded"
    data = (await client.get(root)).json()
    assert data["readiness"]["approved"] is preapproved
    assert data["workspace"]["canonical_values"] == prior_values
    assert data["workspace"]["approved_revision"] == (
        data["workspace"]["revision"] if preapproved else None
    )
