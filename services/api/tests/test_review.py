import base64
import io
from uuid import uuid4

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from ecus_api.core.errors import AppError
from ecus_api.modules.review.domain import evaluate, normalize
from ecus_api.modules.review.schemas import Document, Mapping, Observation, Template, Workspace


def pdf_bytes(lines: list[str]) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    escaped = [line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for line in lines]
    stream.set_data(
        (
            "BT /F1 12 Tf 40 740 Td "
            + " 0 -20 Td ".join(f"({line}) Tj" for line in escaped)
            + " ET"
        ).encode()
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Review"
    sheet["A1"] = "Test template"
    sheet["B2"] = "replace me"
    sheet["B2"].font = Font(bold=True)
    sheet["C2"] = "=B2*2"
    sheet.merge_cells("D2:E2")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload(content: bytes, filename: str, revision: int) -> dict[str, object]:
    return dict(
        content_base64=base64.b64encode(content).decode(), filename=filename, revision=revision
    )


async def new_case(client: httpx.AsyncClient) -> str:
    result = await client.post(
        "/api/v1/cases",
        json=dict(reference=f"R-{uuid4().hex[:8]}", title="Synthetic test", procedure="TEST"),
    )
    assert result.status_code == 201
    return f"/api/v1/cases/{result.json()['id']}/review"


@pytest.mark.parametrize(
    ("raw", "unit", "locale", "expected"),
    [
        ("21", "MT", "en", ("21000", "KG")),
        ("21,000", "KGS", "en", ("21000", "KG")),
        ("21.000,50", "kg", "vi", ("21000.50", "KG")),
        ("20,160.00", "usd", "en", ("20160.00", "USD")),
    ],
)
def test_numeric_normalization(raw: str, unit: str, locale: str, expected: tuple[str, str]) -> None:
    assert normalize(raw, "decimal", unit, locale) == expected


@pytest.mark.parametrize("raw", ["1,2", "NaN", "Infinity", "12,34.00", "1e4", "21 KG"])
def test_ambiguous_numbers_rejected(raw: str) -> None:
    with pytest.raises(AppError, match="specified locale"):
        normalize(raw, "decimal", "KG", "en")


@pytest.mark.parametrize(
    ("values", "verified", "expected"),
    [
        ([], True, "MISSING_REQUIRED"),
        (["500"], True, "NOT_COMPARABLE"),
        (["500", "500.0"], True, "MATCH"),
        (["500", "500", "505"], True, "MISMATCH"),
        (["500", "500", "505", "505"], True, "NEEDS_REVIEW"),
        (["500"], False, "EXTRACTION_UNCERTAIN"),
        (["500", "505"], False, "EXTRACTION_UNCERTAIN"),
        (["500", "500.0"], False, "EXTRACTION_UNCERTAIN"),
    ],
)
def test_comparison_taxonomy(values: list[str], verified: bool, expected: str) -> None:
    state = Workspace(
        template=Template(
            filename="test.xlsx",
            storage_key="key",
            version="v1",
            mapping_version="m1",
            mappings=[Mapping(field="quantity", sheet="Review", cell="B2")],
        )
    )
    for value in values:
        document = Document(
            id=uuid4(), filename="source.pdf", storage_key=uuid4().hex, pages=[value]
        )
        state.documents.append(document)
        state.observations.append(
            Observation(
                id=uuid4(),
                field="quantity",
                document_id=document.id,
                document_type="INVOICE",
                page=1,
                quote=value,
                offset=0,
                normalized=value,
                value_type="decimal",
                unit="KG",
                verified=verified,
            )
        )
    result = evaluate(state)
    from decimal import Decimal

    assert result.comparisons[0].source_values_differ == (len({Decimal(v) for v in values}) > 1)
    assert all(item.verified == verified for item in state.observations)
    if not verified:
        assert not result.eligible

    assert result.comparisons[0].status == expected
    assert not result.eligible  # A required value must be selected, even when sources match.


async def test_real_sources_server_gate_export_and_revision_invalidation(
    client: httpx.AsyncClient,
) -> None:
    root = await new_case(client)
    assert (await client.post(f"{root}/exports", json={"revision": 0})).status_code == 409
    content = pdf_bytes(["Invoice", "21,000", "00123", "=1+1"])
    response = await client.post(f"{root}/documents", json=upload(content, "source.pdf", 0))
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    document = state["documents"][0]
    assert "21,000" in document["pages"][0]
    assert (await client.get(f"{root}/documents/{document['id']}/content")).content == content
    other = await new_case(client)
    assert (await client.get(f"{other}/documents/{document['id']}/content")).status_code == 404
    invalid = dict(
        revision=state["revision"],
        field="quantity",
        document_id=document["id"],
        document_type="INVOICE",
        page=1,
        quote="99999",
        value_type="decimal",
        unit="KG",
    )
    assert (await client.post(f"{root}/observations", json=invalid)).json()["error"][
        "code"
    ] == "source_quote_missing"
    result = await client.post(f"{root}/observations", json={**invalid, "quote": "21,000"})
    state = result.json()["workspace"]
    observation = state["observations"][0]
    assert observation["offset"] == document["pages"][0].index("21,000")
    assert result.json()["readiness"]["comparisons"][0]["status"] == "EXTRACTION_UNCERTAIN"
    selected = dict(
        revision=state["revision"], observation_id=observation["id"], reason="Checked original"
    )
    assert (
        await client.post(f"{root}/fields/quantity/selection", json=selected)
    ).status_code == 422
    verified = await client.post(
        f"{root}/observations/{observation['id']}/verify",
        json=dict(revision=state["revision"], reason="Read original PDF"),
    )
    state = verified.json()["workspace"]
    result = await client.post(
        f"{root}/fields/quantity/selection", json={**selected, "revision": state["revision"]}
    )
    state = result.json()["workspace"]
    result = await client.post(
        f"{root}/policy",
        json=dict(
            revision=state["revision"],
            reason="Synthetic test policy",
            policy=dict(
                version="test-v1",
                name="Test rules",
                status="project_baseline",
                fields=[dict(field="quantity", label="Quantity", value_type="decimal", unit="KG")],
            ),
        ),
    )
    state = result.json()["workspace"]
    result = await client.post(
        f"{root}/logical-documents",
        json=dict(
            revision=state["revision"],
            reason="Read source",
            document_id=document["id"],
            document_type="INVOICE",
            page_from=1,
            page_to=1,
        ),
    )
    state = result.json()["workspace"]
    mappings = [dict(field="quantity", sheet="Review", cell="B2", required=True)]
    template = {
        **upload(template_bytes(), "template.xlsx", state["revision"]),
        "version": "v1",
        "mapping_version": "m1",
        "mappings": mappings,
    }
    result = await client.post(f"{root}/template", json=template)
    assert result.status_code == 200, result.text
    state = result.json()["workspace"]
    assert result.json()["readiness"]["eligible"]
    assert (
        await client.post(f"{root}/exports", json={"revision": state["revision"]})
    ).status_code == 409
    result = await client.post(
        f"{root}/approve",
        json=dict(revision=state["revision"], reason="All required fields checked"),
    )
    state = result.json()["workspace"]
    approved_revision = state["revision"]
    assert result.json()["readiness"]["approved"]
    result = await client.post(f"{root}/exports", json=dict(revision=state["revision"]))
    assert result.status_code == 200, result.text
    state = result.json()["workspace"]
    export = state["exports"][0]
    assert export["source_revision"] == approved_revision
    assert export["policy_snapshot"]["version"] == "test-v1"
    assert export["logical_documents_snapshot"][0]["document_type"] == "INVOICE"
    output = await client.get(f"{root}/exports/{export['id']}/content")
    workbook = load_workbook(io.BytesIO(output.content))
    assert workbook["Review"]["B2"].value == 21000
    assert workbook["Review"]["B2"].font.bold
    assert workbook["Review"]["C2"].value == "=B2*2"
    assert export["template_version"] == "v1" and export["mapping_version"] == "m1"
    assert state["audit"][-1]["action"] == "excel_exported"
    assert (await client.get(f"{other}/exports/{export['id']}/content")).status_code == 404
    assert (
        await client.post(
            f"{root}/approve", json=dict(revision=approved_revision, reason="Stale tab")
        )
    ).status_code == 409
    result = await client.post(
        f"{root}/fields/quantity/selection",
        json={**selected, "revision": state["revision"], "reason": "Changed review decision"},
    )
    state = result.json()["workspace"]
    assert not result.json()["readiness"]["approved"]
    assert (
        await client.post(f"{root}/exports", json={"revision": state["revision"]})
    ).status_code == 409
    assert (await client.get(root)).json()["workspace"] == state


@pytest.mark.parametrize("cell", ["C2", "E2", "XFE2", "A1048577"])
async def test_invalid_excel_targets_rejected(client: httpx.AsyncClient, cell: str) -> None:
    root = await new_case(client)
    data = {
        **upload(template_bytes(), "test.xlsx", 0),
        "version": "v1",
        "mapping_version": "m1",
        "mappings": [dict(field="value", sheet="Review", cell=cell)],
    }
    result = await client.post(f"{root}/template", json=data)
    assert result.status_code == 422
    assert result.json()["error"]["code"] == "invalid_mapping"


async def test_duplicate_files_missing_fields_and_unreadable_source_block(
    client: httpx.AsyncClient,
) -> None:
    root = await new_case(client)
    file = pdf_bytes([])
    result = await client.post(f"{root}/documents", json=upload(file, "scan.pdf", 0))
    assert result.status_code == 200
    assert "scan.pdf:source_unreadable" in result.json()["readiness"]["blockers"]
    state = result.json()["workspace"]
    duplicate = await client.post(
        f"{root}/documents", json=upload(file, "renamed.pdf", state["revision"])
    )
    assert duplicate.status_code == 409
    template = {
        **upload(template_bytes(), "test.xlsx", state["revision"]),
        "version": "v1",
        "mapping_version": "m1",
        "mappings": [dict(field="gross_weight", sheet="Review", cell="B2")],
    }
    result = await client.post(f"{root}/template", json=template)
    assert result.json()["readiness"]["comparisons"][0]["status"] == "MISSING_REQUIRED"
    assert (
        await client.post(
            f"{root}/approve",
            json=dict(revision=result.json()["workspace"]["revision"], reason="Cannot bypass gate"),
        )
    ).status_code == 409


@pytest.mark.parametrize("value", ["00123", "=1+1", "+SUM(A1:A2)"])
def test_excel_identifier_and_formula_like_text_stay_literal(value: str) -> None:
    from ecus_api.modules.review.schemas import Selection
    from ecus_api.modules.review.service import generate_excel

    document_id = uuid4()
    observation = Observation(
        id=uuid4(),
        field="identifier",
        document_id=document_id,
        document_type="INVOICE",
        page=1,
        quote=value,
        offset=0,
        normalized=value,
        value_type="text",
        unit="",
        verified=True,
    )
    state = Workspace(
        observations=[observation],
        selections={
            "identifier": Selection(observation_id=observation.id, reason="Source checked")
        },
        template=Template(
            filename="test.xlsx",
            storage_key="key",
            version="v1",
            mapping_version="m1",
            mappings=[Mapping(field="identifier", sheet="Review", cell="B2")],
        ),
    )
    output = load_workbook(io.BytesIO(generate_excel(state, template_bytes())))
    assert output["Review"]["B2"].value == value
    assert output["Review"]["B2"].data_type == "s"


async def test_mapping_versions_cannot_be_redefined(client: httpx.AsyncClient) -> None:
    root = await new_case(client)
    original = {
        **upload(template_bytes(), "template.xlsx", 0),
        "version": "v1",
        "mapping_version": "m1",
        "mappings": [dict(field="quantity", sheet="Review", cell="B2")],
    }
    result = await client.post(f"{root}/template", json=original)
    assert result.status_code == 200
    revision = result.json()["workspace"]["revision"]
    changed = {
        **original,
        "revision": revision,
        "mappings": [dict(field="quantity", sheet="Review", cell="B3")],
    }
    conflict = await client.post(f"{root}/template", json=changed)
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "version_conflict"
    accepted = await client.post(f"{root}/template", json={**changed, "mapping_version": "m2"})
    assert accepted.status_code == 200
    assert len(accepted.json()["workspace"]["template_history"]) == 2


async def test_two_tabs_cannot_commit_the_same_revision(client: httpx.AsyncClient) -> None:
    import asyncio

    root = await new_case(client)
    results = await asyncio.gather(
        *[
            client.post(f"{root}/documents", json=upload(pdf_bytes([text]), f"{text}.pdf", 0))
            for text in ["One", "Two"]
        ]
    )
    assert sorted(response.status_code for response in results) == [200, 409]
    state = (await client.get(root)).json()["workspace"]
    assert state["revision"] == 1
    assert len(state["documents"]) == 1
    assert len(state["audit"]) == 1
