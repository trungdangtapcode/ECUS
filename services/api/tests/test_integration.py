"""End-to-end preparation tests use synthetic documents and blank layouts only."""

import io
from pathlib import Path

import httpx
import pytest
import xlwt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from test_review import new_case, pdf_bytes, upload

from ecus_api.core.config import Settings
from ecus_api.core.errors import AppError
from ecus_api.db.session import Database
from ecus_api.modules.review.legacy_excel import blank_xls_to_xlsx
from ecus_api.modules.review.template_library import Layout, TemplateProfile, expand_layout
from ecus_api.storage.local import LocalBlobStore


def synthetic_layout() -> Layout:
    return Layout.model_validate(
        {
            "id": "synthetic-form-v1",
            "name_key": "integration.exportForm",
            "version": "synthetic-v1",
            "caption_cells": {"Form": ["A1"], "Goods": ["A1"]},
            "fields": [
                {
                    "field": "quantity",
                    "label": "Quantity",
                    "value_type": "decimal",
                    "unit": "KG",
                    "targets": [{"sheet": "Form", "cell": "B2"}],
                },
                {
                    "field": "identifier",
                    "label": "Identifier",
                    "targets": [{"sheet": "Form", "cell": "B3"}, {"sheet": "Goods", "cell": "B5"}],
                },
                {
                    "field": "missing",
                    "label": "Missing",
                    "targets": [{"sheet": "Form", "cell": "D2"}],
                },
            ],
            "groups": [
                {
                    "key": "goods",
                    "maximum": 5,
                    "fields": [
                        {
                            "field": "description",
                            "label": "Description",
                            "targets": [{"sheet": "Goods", "cell": "B2"}],
                        }
                    ],
                    "slots": [{"sheet": "Goods"}],
                    "repeat_sheet": "Goods",
                    "repeat_offsets": [0],
                }
            ],
        }
    )


def blank_form() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Form"
    sheet["A1"] = "Synthetic form"
    sheet["B2"].font = Font(bold=True)
    sheet.merge_cells("D2:E2")
    sheet["C4"] = "=B2*2"
    workbook.create_sheet("Goods")["A1"] = "Synthetic goods"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def register_profile(database: Database, settings: Settings) -> None:
    layout = synthetic_layout()
    key = LocalBlobStore(settings.storage_root).put(blank_form())
    async with database.sessions() as session:
        session.add(
            TemplateProfile(
                id=layout.id,
                definition=layout.model_dump(mode="json"),
                storage_key=key,
                enabled=True,
            )
        )
        await session.commit()


async def test_full_profile_conflict_edit_missing_red_export_and_stale_sources(
    client: httpx.AsyncClient, database: Database, settings: Settings
) -> None:
    await register_profile(database, settings)
    root = await new_case(client)
    profiles = (await client.get("/api/v1/reference/excel-templates")).json()
    assert profiles[0]["groups"] == {"goods": 5}
    response = await client.post(
        f"{root}/template-profile",
        json={
            "revision": 0,
            "profile_id": profiles[0]["id"],
            "counts": {"goods": 4},
            "reason": "Selected synthetic blank form",
        },
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    for name, text in [("first.pdf", "100"), ("second.pdf", "101")]:
        response = await client.post(
            f"{root}/documents", json=upload(pdf_bytes([text]), name, state["revision"])
        )
        state = response.json()["workspace"]
        document = state["documents"][-1]
        response = await client.post(
            f"{root}/observations",
            json={
                "revision": state["revision"],
                "field": "quantity",
                "document_id": document["id"],
                "document_type": "INVOICE",
                "page": 1,
                "quote": text,
                "value_type": "decimal",
                "unit": "KG",
                "verified": True,
            },
        )
        state = response.json()["workspace"]
    assert not response.json()["readiness"]["eligible"]
    assert "quantity:selection_required" in response.json()["readiness"]["blockers"]
    assert "missing:missing_value" in response.json()["readiness"]["warnings"]
    for field, value in [
        ("quantity", "100.50"),
        ("identifier", "00123"),
        ("goods_004_description", "=1+1"),
    ]:
        response = await client.post(
            f"{root}/fields/{field}/canonical",
            json={
                "revision": state["revision"],
                "value": value,
                "reason": "Reviewed original and entered corrected output",
                "reviewed": True,
            },
        )
        assert response.status_code == 200, response.text
        state = response.json()["workspace"]
    readiness = response.json()["readiness"]
    quantity = next(item for item in readiness["comparisons"] if item["field"] == "quantity")
    assert quantity["status"] == "NEEDS_REVIEW" and quantity["resolved"]
    assert [item["quote"] for item in quantity["observations"]] == ["100", "101"]
    assert readiness["eligible"]
    # Same profile can be selected again without redefining its immutable workbook version.
    response = await client.post(
        f"{root}/template-profile",
        json={
            "revision": state["revision"],
            "profile_id": profiles[0]["id"],
            "counts": {"goods": 4},
            "reason": "Confirmed counts",
        },
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    smaller = await client.post(
        f"{root}/template-profile",
        json={
            "revision": state["revision"],
            "profile_id": profiles[0]["id"],
            "counts": {"goods": 1},
            "reason": "Would remove reviewed goods",
        },
    )
    assert smaller.status_code == 409
    response = await client.post(
        f"{root}/approve",
        json={"revision": state["revision"], "reason": "Accept missing fields marked red"},
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    response = await client.post(f"{root}/exports", json={"revision": state["revision"]})
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    record = state["exports"][0]
    assert record["canonical_snapshot"]["quantity"]["value"] == "100.50"
    assert record["values_snapshot"]["identifier"] == "00123"
    assert "missing" in record["missing_fields"]
    output = await client.get(f"{root}/exports/{record['id']}/content")
    book = load_workbook(io.BytesIO(output.content))
    assert book["Form"]["B2"].value == 100.5 and book["Form"]["B2"].font.bold
    assert book["Form"]["B3"].value == "00123"
    assert book["Goods_04"]["B2"].value == "=1+1" and book["Goods_04"]["B2"].data_type == "s"
    assert book["Goods_04"]["B5"].value == "00123"
    assert book["Form"]["C4"].value == "=B2*2"
    assert book["Form"]["D2"].value is None and book["Form"]["D2"].fill.fgColor.rgb == "FFF2CCCC"
    response = await client.post(
        f"{root}/documents", json=upload(pdf_bytes(["new evidence"]), "new.pdf", state["revision"])
    )
    state = response.json()["workspace"]
    assert not response.json()["readiness"]["approved"]
    assert "quantity:canonical_edit_stale" in response.json()["readiness"]["blockers"]
    assert (
        await client.post(f"{root}/exports", json={"revision": state["revision"]})
    ).status_code == 409
    assert (await client.get(f"{root}/exports/{record['id']}/content")).content == output.content
    stale = await client.post(
        f"{root}/fields/quantity/canonical",
        json={"revision": state["revision"] - 1, "value": "99", "reason": "Old browser"},
    )
    assert stale.status_code == 409
    response = await client.post(
        f"{root}/fields/quantity/canonical",
        json={
            "revision": state["revision"],
            "value": "100.50",
            "reason": "Rechecked new evidence",
            "reviewed": True,
        },
    )
    state = response.json()["workspace"]
    assert state["canonical_history"][0]["quantity"]["value"] == "100.50"
    response = await client.post(
        f"{root}/fields/quantity/canonical/withdraw",
        json={"revision": state["revision"], "reason": "Return to evidence selection"},
    )
    assert "quantity:selection_required" in response.json()["readiness"]["blockers"]


async def test_manual_edit_cannot_suppress_unverified_ai_values(
    client: httpx.AsyncClient, database: Database, settings: Settings
) -> None:
    await register_profile(database, settings)
    root = await new_case(client)
    response = await client.post(
        f"{root}/template-profile",
        json={
            "revision": 0,
            "profile_id": "synthetic-form-v1",
            "counts": {"goods": 0},
            "reason": "Selected form",
        },
    )
    state = response.json()["workspace"]
    response = await client.post(
        f"{root}/documents", json=upload(pdf_bytes(["99"]), "source.pdf", state["revision"])
    )
    state = response.json()["workspace"]
    response = await client.post(
        f"{root}/observations",
        json={
            "revision": state["revision"],
            "field": "quantity",
            "document_id": state["documents"][0]["id"],
            "document_type": "INVOICE",
            "page": 1,
            "quote": "99",
            "value_type": "decimal",
            "unit": "KG",
        },
    )
    state = response.json()["workspace"]
    response = await client.post(
        f"{root}/fields/quantity/canonical",
        json={"revision": state["revision"], "value": "99", "reason": "Manual output edit"},
    )
    state = response.json()["workspace"]
    assert "quantity:EXTRACTION_UNCERTAIN" in response.json()["readiness"]["blockers"]
    assert (
        await client.post(
            f"{root}/approve",
            json={"revision": state["revision"], "reason": "Attempt to skip verification"},
        )
    ).status_code == 409
    for value in ["NaN", "1e4", "   "]:
        result = await client.post(
            f"{root}/fields/quantity/canonical",
            json={"revision": state["revision"], "value": value, "reason": "Invalid decimal"},
        )
        assert result.status_code == 422


async def test_unconfirmed_manual_value_blocks_missing_and_verified_sources(
    client: httpx.AsyncClient, database: Database, settings: Settings
) -> None:
    await register_profile(database, settings)
    root = await new_case(client)
    state = {"revision": 0}

    async def post(path: str, body: dict) -> dict:
        nonlocal state
        response = await client.post(root + path, json={"revision": state["revision"], **body})
        assert response.status_code == 200, response.text
        data = response.json()
        state = data["workspace"]
        return data

    await post(
        "/template-profile",
        {"profile_id": "synthetic-form-v1", "counts": {"goods": 0}, "reason": "Selected form"},
    )
    await post("/documents", upload(pdf_bytes(["100"]), "source.pdf", state["revision"]))
    await post(
        "/observations",
        {
            "field": "quantity",
            "document_id": state["documents"][0]["id"],
            "document_type": "INVOICE",
            "page": 1,
            "quote": "100",
            "value_type": "decimal",
            "unit": "KG",
            "verified": True,
        },
    )
    for field, value in [("quantity", "100.50"), ("identifier", "00123")]:
        data = await post(f"/fields/{field}/canonical", {"value": value, "reason": "Draft edit"})
        assert f"{field}:canonical_review_required" in data["readiness"]["blockers"]
        assert (
            await client.post(
                root + "/approve", json={"revision": state["revision"], "reason": "Attempt"}
            )
        ).status_code == 409
        assert (
            await client.post(root + "/exports", json={"revision": state["revision"]})
        ).status_code == 409
        await post(
            f"/fields/{field}/canonical",
            {"value": value, "reason": "Checked source and corrected output", "reviewed": True},
        )
    assert (await client.get(root)).json()["readiness"]["eligible"]


async def test_verified_transcript_preserves_original_pdf_and_anchors_observation(
    client: httpx.AsyncClient,
) -> None:
    root = await new_case(client)
    original = pdf_bytes([])
    response = await client.post(f"{root}/documents", json=upload(original, "scan.pdf", 0))
    state = response.json()["workspace"]
    document = state["documents"][0]
    response = await client.post(
        f"{root}/documents/{document['id']}/transcript",
        json={
            "revision": state["revision"],
            "page": 1,
            "text": "Invoice\n00123",
            "reason": "Manually checked original page",
        },
    )
    assert response.status_code == 200
    state = response.json()["workspace"]
    assert state["documents"][0]["pages"] == [""]
    assert "scan.pdf:source_unreadable" not in response.json()["readiness"]["blockers"]
    assert (await client.get(f"{root}/documents/{document['id']}/content")).content == original
    response = await client.post(
        f"{root}/observations",
        json={
            "revision": state["revision"],
            "field": "identifier",
            "document_id": document["id"],
            "document_type": "INVOICE",
            "page": 1,
            "quote": "00123",
        },
    )
    assert response.status_code == 200
    state = response.json()["workspace"]
    assert state["observations"][0]["offset"] == 8
    response = await client.post(
        f"{root}/documents/{document['id']}/transcript",
        json={
            "revision": state["revision"],
            "page": 1,
            "text": "changed",
            "reason": "Cannot replace cited text",
        },
    )
    assert response.status_code == 409


@pytest.mark.parametrize("count", [-1, 6])
def test_repeated_records_never_truncate(count: int) -> None:
    with pytest.raises(AppError):
        expand_layout(blank_form(), synthetic_layout(), {"goods": count})


def test_legacy_form_sanitization_and_formula_rejection(tmp_path: Path) -> None:
    source = xlwt.Workbook()
    sheet = source.add_sheet("Form")
    sheet.col(1).width = 4000
    sheet.write(0, 0, "Caption", xlwt.easyxf("font: bold on;"))
    sheet.write(1, 1, "OLD SHIPMENT VALUE")
    sheet.write_merge(1, 1, 3, 4, "OLD MERGED VALUE")
    source.add_sheet("Goods").write(0, 0, "Goods")
    stream = io.BytesIO()
    source.save(stream)
    result = blank_xls_to_xlsx(stream.getvalue(), synthetic_layout())
    workbook = load_workbook(io.BytesIO(result))
    assert workbook["Form"]["A1"].value == "Caption" and workbook["Form"]["A1"].font.bold
    assert workbook["Form"]["B2"].value is None and workbook["Form"]["D2"].value is None
    assert "D2:E2" in workbook["Form"].merged_cells
    assert workbook["Form"].column_dimensions["B"].width == 4000 / 256
    sheet.write(2, 0, xlwt.Formula("1+1"))
    stream = io.BytesIO()
    source.save(stream)
    with pytest.raises(AppError) as error:
        blank_xls_to_xlsx(stream.getvalue(), synthetic_layout())
    assert error.value.code == "unsupported_legacy_template"


async def test_form_selection_preserves_case_document_policy(
    client: httpx.AsyncClient, database: Database, settings: Settings
) -> None:
    await register_profile(database, settings)
    root = await new_case(client)
    response = await client.post(
        f"{root}/policy",
        json={
            "revision": 0,
            "reason": "Explicit case rules",
            "policy": {
                "version": "governed-v1",
                "name": "Governed document requirements",
                "status": "proposal",
                "document_requirements": [
                    {"key": "invoice", "label": "Invoice", "any_of": ["INVOICE"]}
                ],
                "fields": [
                    {
                        "field": "quantity",
                        "label": "Reviewed quantity",
                        "value_type": "decimal",
                        "unit": "KG",
                        "required": True,
                        "preferred_sources": ["INVOICE"],
                    }
                ],
            },
        },
    )
    assert response.status_code == 200
    state = response.json()["workspace"]
    response = await client.post(
        f"{root}/template-profile",
        json={
            "revision": state["revision"],
            "profile_id": "synthetic-form-v1",
            "counts": {"goods": 1},
            "reason": "Choose layout without changing rules",
        },
    )
    assert response.status_code == 200, response.text
    state = response.json()["workspace"]
    assert state["policy"]["status"] == "proposal"
    assert state["policy"]["document_requirements"][0]["key"] == "invoice"
    quantity = next(rule for rule in state["policy"]["fields"] if rule["field"] == "quantity")
    assert quantity["preferred_sources"] == ["INVOICE"] and quantity["required"]
    assert "policy_proposal" in response.json()["readiness"]["blockers"]
    assert "invoice:logical_document_missing" in response.json()["readiness"]["blockers"]


async def test_repeated_source_text_requires_exact_occurrence(client: httpx.AsyncClient) -> None:
    root = await new_case(client)
    response = await client.post(
        f"{root}/documents", json=upload(pdf_bytes(["Total 100", "Price 100"]), "repeated.pdf", 0)
    )
    state = response.json()["workspace"]
    document = state["documents"][0]
    body = {
        "revision": state["revision"],
        "field": "quantity",
        "document_id": document["id"],
        "document_type": "INVOICE",
        "page": 1,
        "quote": "100",
        "value_type": "decimal",
    }
    response = await client.post(f"{root}/observations", json=body)
    assert (
        response.status_code == 422 and response.json()["error"]["code"] == "source_quote_ambiguous"
    )
    offset = document["pages"][0].rfind("100")
    response = await client.post(f"{root}/observations", json={**body, "source_offset": offset})
    assert response.status_code == 200
    assert response.json()["workspace"]["observations"][0]["offset"] == offset
    state = response.json()["workspace"]
    response = await client.post(
        f"{root}/observations",
        json={**body, "revision": state["revision"], "source_offset": offset - 1},
    )
    assert (
        response.status_code == 422 and response.json()["error"]["code"] == "source_quote_missing"
    )
