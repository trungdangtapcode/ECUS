"""Draft exports never approve unresolved inputs or silently pick an AI proposal."""

import io

import httpx
import pytest
import xlrd
from openpyxl import load_workbook
from test_native_excel import native_form, native_layout
from test_review import new_case, pdf_bytes, template_bytes, upload

from ecus_api.modules.review.native_excel import edit_native


@pytest.mark.parametrize("file_format", ["xls", "xlsx"])
@pytest.mark.parametrize("remark", [None, "Verified supplier remark"])
async def test_draft_preserves_unapproved_gate_and_immutable_review_snapshot(
    client: httpx.AsyncClient,
    file_format: str,
    remark: str | None,
) -> None:
    root = await new_case(client)
    state = (await client.get(root)).json()["workspace"]
    response = await client.post(
        f"{root}/policy",
        json={
            "revision": state["revision"],
            "reason": "Synthetic draft review",
            "policy": {
                "version": "draft-v1",
                "name": "Draft review",
                "status": "project_baseline",
                "fields": [
                    {
                        "field": "quantity",
                        "label": "Quantity",
                        "value_type": "decimal",
                        "unit": "KG",
                    },
                    {"field": "identifier", "label": "Identifier", "required": False},
                    {"field": "conflict", "label": "Conflicting source", "required": False},
                    {"field": "notes", "label": "Notes", "required": False},
                ],
            },
        },
    )
    assert response.status_code == 200
    state = response.json()["workspace"]
    if file_format == "xls":
        content = edit_native(native_form(), {}, captions=native_layout().caption_cells)
        targets = [("Form", "B2"), ("Form", "B3"), ("Form", "D2"), ("Goods", "B5")]
    else:
        content = template_bytes()
        targets = [("Review", "B2"), ("Review", "B3"), ("Review", "D2"), ("Review", "A4")]
    response = await client.post(
        f"{root}/template",
        json={
            **upload(content, f"synthetic.{file_format}", state["revision"]),
            "version": "draft-layout-v1",
            "mapping_version": "draft-mapping-v1",
            "mappings": [
                {"field": field, "sheet": sheet, "cell": cell, "required": False}
                for field, (sheet, cell) in zip(
                    ["quantity", "identifier", "conflict", "notes"], targets, strict=True
                )
            ],
        },
    )
    assert response.status_code == 200
    state = response.json()["workspace"]
    for quote in ["SOURCE_A", "SOURCE_B"]:
        response = await client.post(
            f"{root}/documents",
            json=upload(pdf_bytes(["Invoice", quote]), f"{quote}.pdf", state["revision"]),
        )
        state = response.json()["workspace"]
        document = state["documents"][-1]
        response = await client.post(
            f"{root}/observations",
            json={
                "revision": state["revision"],
                "field": "conflict",
                "document_id": document["id"],
                "document_type": "INVOICE",
                "page": 1,
                "quote": quote,
                "verified": False,
            },
        )
        assert response.status_code == 200
        state = response.json()["workspace"]
    for field, value, reviewed in [
        ("quantity", "21000", True),
        ("identifier", "UNREVIEWED", False),
    ]:
        response = await client.post(
            f"{root}/fields/{field}/canonical",
            json={
                "revision": state["revision"],
                "value": value,
                "reason": "Synthetic operator decision",
                "reviewed": reviewed,
            },
        )
        assert response.status_code == 200
        state = response.json()["workspace"]
    if remark is not None:
        response = await client.post(
            f"{root}/fields/notes/canonical",
            json={
                "revision": state["revision"],
                "value": remark,
                "reason": "Reviewed real remarks",
                "reviewed": True,
            },
        )
        assert response.status_code == 200
        state = response.json()["workspace"]
    original_sources = state["observations"]
    revision = state["revision"]
    response = await client.post(
        f"{root}/draft-exports",
        json={
            "revision": revision,
            "reason": "Send for client review",
            "locale": "en",
        },
    )
    assert response.status_code == 200
    data = response.json()
    state = data["workspace"]
    assert not data["readiness"]["approved"] and not data["readiness"]["eligible"]
    assert state["approved_revision"] is None
    assert state["observations"] == original_sources
    record = state["exports"][-1]
    assert record["kind"] == "draft" and record["filename"].startswith("draft-review-")
    assert record["source_revision"] == revision and record["blockers_snapshot"]
    assert record["annotation_fields"] == []
    assert record["values_snapshot"]["notes"] == remark
    assert record["values_snapshot"]["quantity"] == "21000"
    assert record["values_snapshot"]["identifier"] is None
    assert record["values_snapshot"]["conflict"] is None
    assert record["selected_observations"] == {}
    assert (
        await client.post(
            f"{root}/draft-exports",
            json={
                "revision": revision,
                "reason": "Stale revision",
                "locale": "en",
            },
        )
    ).status_code == 409
    assert (
        await client.post(f"{root}/exports", json={"revision": state["revision"]})
    ).status_code == 409
    assert (
        await client.post(
            f"{root}/approve",
            json={
                "revision": state["revision"],
                "reason": "Still unresolved",
            },
        )
    ).status_code == 409
    response = await client.get(f"{root}/exports/{record['id']}/content")
    assert response.status_code == 200
    exported = response.content
    if file_format == "xls":
        book = xlrd.open_workbook(file_contents=exported, formatting_info=True)
        form = book.sheet_by_name("Form")
        assert form.cell_value(1, 1) == 21000
        assert form.cell_value(2, 1) == "" and form.cell_value(1, 3) == ""
        for row, column in [(2, 1), (1, 3), (1, 4)]:
            xf = book.xf_list[form.cell(row, column).xf_index]
            assert book.colour_map[xf.background.pattern_colour_index] == (242, 204, 204)
        notice = book.sheet_by_name("Goods").cell_value(4, 1)
    else:
        book = load_workbook(io.BytesIO(exported))
        form = book["Review"]
        assert form["B2"].value == 21000
        assert form["B3"].value is None and form["D2"].value is None
        assert form["D2"].fill.fgColor.rgb == "FFF2CCCC"
        assert form["C2"].value == "=B2*2"
        notice = form["A4"].value
    assert (notice or None) == remark
    assert "DRAFT" not in (notice or "")
    if remark is None:
        assert state["canonical_values"].get("notes") is None
    response = await client.post(
        f"{root}/fields/quantity/canonical",
        json={
            "revision": state["revision"],
            "value": "22000",
            "reason": "Later revision",
            "reviewed": True,
        },
    )
    assert response.status_code == 200
    assert (await client.get(f"{root}/exports/{record['id']}/content")).content == exported
    assert response.json()["workspace"]["exports"][-1]["values_snapshot"]["quantity"] == "21000"
