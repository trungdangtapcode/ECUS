"""Create synthetic blank forms only inside the isolated browser-test environment."""

import asyncio
import base64
import io
from pathlib import Path

import xlwt
from ecus_api.core.config import Settings
from ecus_api.db.session import Database
from ecus_api.modules.review.schemas import TemplateInput
from ecus_api.modules.review.service import validate_template
from ecus_api.modules.review.template_library import Layout, TemplateProfile, expand_layout
from ecus_api.storage.local import LocalBlobStore
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

PROFILE_IDS = {
    "export-form-v1",
    "export-form-v2",
    "original-form-xls-v6",
    "original-form-xls-v8",
    "original-form-xls-v9",
}


def synthetic_form(layout: Layout) -> bytes:
    cells: dict[str, dict[str, str]] = {}

    def cell(sheet: str, address: str, text: str = "") -> None:
        cells.setdefault(sheet, {}).setdefault(address, text)

    for sheet, captions in layout.caption_cells.items():
        for address in captions:
            cell(sheet, address, "Synthetic test heading")
    for sheet, headings in layout.headings.items():
        for address in headings:
            cell(sheet, address, "Synthetic test heading")
    for field in layout.fields:
        for target in field.targets:
            cell(target.sheet, target.cell)
    for group in layout.groups:
        if group.repeat_sheet:
            cells.setdefault(group.repeat_sheet, {})
        for slot in group.slots:
            cells.setdefault(slot.sheet, {})
            for field in group.fields:
                overrides = slot.field_targets.get(field.field)
                for target in overrides or field.targets:
                    row, column = coordinate_to_tuple(target.cell)
                    address = (
                        target.cell
                        if overrides
                        else f"{get_column_letter(column + slot.column_offset)}"
                        f"{row + slot.row_offset}"
                    )
                    cell(target.sheet if overrides else slot.sheet, address)
    output = io.BytesIO()
    if layout.file_format == "xls":
        native = xlwt.Workbook()
        style = xlwt.easyxf("alignment: wrap on")
        for name, content in cells.items():
            sheet = native.add_sheet(name)
            for address, text in content.items():
                row, column = coordinate_to_tuple(address)
                sheet.write(row - 1, column - 1, text, style)
        native.save(output)
    else:
        workbook = Workbook()
        workbook.remove(workbook[workbook.sheetnames[0]])
        for name, content in cells.items():
            sheet = workbook.create_sheet(name)
            for address, text in content.items():
                sheet[address] = text
                sheet[address].alignment = Alignment(wrap_text=True)
        workbook.save(output)
        workbook.close()
    return output.getvalue()


async def seed_templates(settings: Settings, directory: Path) -> None:
    if settings.app_env != "test" or not settings.ai_test_transport:
        raise RuntimeError("Synthetic forms require the isolated test environment and AI stub.")
    database = Database(settings)
    store = LocalBlobStore(settings.storage_root)
    try:
        async with database.sessions() as session:
            paths = await asyncio.to_thread(lambda: sorted(directory.glob("*.json")))
            for path in paths:
                layout = Layout.model_validate_json(await asyncio.to_thread(path.read_text))
                if layout.id not in PROFILE_IDS:
                    continue
                if await session.get(TemplateProfile, layout.id):
                    continue
                content = synthetic_form(layout)
                expanded, mappings, _ = expand_layout(
                    content, layout, {group.key: group.maximum for group in layout.groups}
                )
                validate_template(
                    expanded,
                    TemplateInput(
                        revision=0,
                        filename=f"synthetic.{layout.file_format}",
                        content_base64=base64.b64encode(expanded).decode(),
                        version=layout.version,
                        mapping_version=layout.version,
                        mappings=mappings,
                    ),
                    settings.review_max_file_bytes,
                )
                session.add(
                    TemplateProfile(
                        id=layout.id,
                        definition=layout.model_dump(mode="json"),
                        storage_key=store.put(content),
                        enabled=True,
                    )
                )
            await session.commit()
        print("Synthetic browser-test forms initialized without private source documents.")
    finally:
        await database.close()


if __name__ == "__main__":
    asyncio.run(seed_templates(Settings(), Path("reference-data/excel-layouts")))  # type: ignore[call-arg]
